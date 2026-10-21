"""Self-contained H3 calibration consistency report generator.

The module intentionally avoids pandas so the packaged monitor can generate a
formatted XLSX report without requiring a separate Python or Node.js install.
"""

from __future__ import annotations

import csv
import math
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from statistics import mean, median, stdev
from typing import Callable, Iterable

import numpy as np
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


FRAME_DIAGNOSTICS = ["upper_model_load_score", "upper_model_soft_gate_rescue"]
SUMMARY_DIAGNOSTICS = ["upper_model_load_score_median", "upper_model_soft_gate_rescue_frames"]
POSITION_ORDER = {"前": 0, "中": 1, "后": 2, "front": 0, "middle": 1, "rear": 2}
LOAD_WORDS = {"加载", "load", "loading"}
UNLOAD_WORDS = {"卸载", "unload", "unloading"}

NAVY = "17365D"
BLUE = "1F4E78"
CYAN = "DDEBF7"
PALE = "EAF3F8"
GREEN = "E2F0D9"
RED = "FCE4D6"
YELLOW = "FFF2CC"
WHITE = "FFFFFF"
GRID = "B8C7D1"


def _text(value: object, default: str = "") -> str:
    return str(value).strip() if value is not None and str(value).strip() else default


def _number(value: object, default: float = math.nan) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return default
    return result if math.isfinite(result) else default


def _percentile(values: Iterable[float], percentile: float) -> float | None:
    finite = [float(value) for value in values if math.isfinite(float(value))]
    return float(np.percentile(finite, percentile)) if finite else None


def _median(values: Iterable[float]) -> float | None:
    finite = [float(value) for value in values if math.isfinite(float(value))]
    return float(median(finite)) if finite else None


def _mean(values: Iterable[float]) -> float | None:
    finite = [float(value) for value in values if math.isfinite(float(value))]
    return float(mean(finite)) if finite else None


def _std(values: Iterable[float]) -> float | None:
    finite = [float(value) for value in values if math.isfinite(float(value))]
    return float(stdev(finite)) if len(finite) >= 2 else None


def _direction(value: object) -> str:
    value = _text(value)
    lowered = value.lower()
    if lowered in LOAD_WORDS:
        return "加载"
    if lowered in UNLOAD_WORDS:
        return "卸载"
    return value


def _session_id(path: Path) -> str:
    match = re.search(r"(\d{8}_\d{6})", path.stem)
    return match.group(1) if match else path.stem.replace("_summary", "")


def read_recoverable_csv(path: Path, kind: str) -> tuple[list[dict[str, str]], bool]:
    """Read a CSV and repair the two missing v5.2.1 diagnostic headings."""
    with path.open("r", newline="", encoding="utf-8-sig") as stream:
        rows = list(csv.reader(stream))
    if not rows:
        return [], False
    header = list(rows[0])
    data = rows[1:]
    first = next((row for row in data if row), None)
    recovered = False
    if first and len(first) == len(header) + 2:
        marker = "upper_model_confidence" if kind == "frames" else "upper_model_median_n"
        additions = FRAME_DIAGNOSTICS if kind == "frames" else SUMMARY_DIAGNOSTICS
        if marker not in header:
            raise ValueError(f"{path.name} 缺少关键列 {marker}，无法修复旧表头。")
        insert_at = header.index(marker) + 1
        header[insert_at:insert_at] = additions
        recovered = True
    parsed: list[dict[str, str]] = []
    for line_number, row in enumerate(data, start=2):
        if not row or not any(cell.strip() for cell in row):
            continue
        if len(row) != len(header):
            raise ValueError(
                f"{path.name} 第 {line_number} 行字段数为 {len(row)}，表头为 {len(header)}。"
            )
        parsed.append(dict(zip(header, row)))
    return parsed, recovered


def _frames_metadata(path: Path) -> tuple[dict[str, str], int, bool]:
    if not path.exists():
        return {}, 0, False
    rows, recovered = read_recoverable_csv(path, "frames")
    metadata: dict[str, str] = {}
    for name in ("device_sn", "silicone_batch", "magnet_assembly", "operator"):
        for row in rows:
            value = _text(row.get(name))
            if value:
                metadata[name] = value
                break
    return metadata, len(rows), recovered


def load_sessions(
    input_dir: str | Path,
    progress: Callable[[int, str], None] | None = None,
) -> tuple[list[dict[str, object]], list[dict[str, object]], dict[str, dict[str, str]]]:
    root = Path(str(input_dir).strip().strip('"'))
    if not root.exists():
        raise ValueError(f"数据路径不存在：{root}")
    paths = ([root] if root.is_file() and root.suffix.lower() == ".csv" else sorted(
        p for p in root.rglob("*") if p.is_file() and
        (p.name.lower() == "summary.csv" or p.name.lower().endswith("_summary.csv"))
    ))
    if not paths:
        raise ValueError(
            f"此路径没有标定 summary CSV：{root}\n"
            "如果是“记录与导出”的普通 CSV，请将数据类型切换为“记录 CSV：前后六工况”。"
        )
    all_rows: list[dict[str, object]] = []
    sources: list[dict[str, object]] = []
    metadata_by_device: dict[str, dict[str, str]] = {}
    for index, summary_path in enumerate(paths, start=1):
        if progress:
            progress(5 + int(index / max(len(paths), 1) * 30), f"读取 {summary_path.name}")
        summary, recovered = read_recoverable_csv(summary_path, "summary")
        frames_name = "frames.csv" if summary_path.name.lower() == "summary.csv" else re.sub(
            r"_summary\.csv$", "_frames.csv", summary_path.name, flags=re.I
        )
        frames_path = summary_path.with_name(frames_name)
        if frames_path == summary_path:
            frames_path = summary_path.with_name(summary_path.stem + "_frames.csv")
        metadata, frame_rows, frames_recovered = _frames_metadata(frames_path)
        session_id = _session_id(summary_path)
        source = {
            "session_id": session_id,
            "summary_file": str(summary_path),
            "frames_file": str(frames_path) if frames_path.exists() else "",
            "summary_rows": len(summary),
            "frame_rows": frame_rows,
            "header_recovered": recovered or frames_recovered,
            "status": "可用" if summary else "空文件，已跳过",
        }
        sources.append(source)
        if not summary:
            continue
        required = {"position", "direction", "repeat", "step_index", "reference_median_n", "upper_model_median_n"}
        missing = required - summary[0].keys()
        if missing:
            raise ValueError(
                f"{summary_path.name} 不是标定汇总格式，缺少列：{', '.join(sorted(missing))}。\n"
                "普通记录 CSV 请使用“记录 CSV：前后六工况”；不要通过重命名转换格式。"
            )
        fallback_device = metadata.get("device_sn") or summary_path.parent.name or session_id
        for line_number, row in enumerate(summary, 2):
            for key in ("repeat", "step_index", "reference_median_n", "upper_model_median_n"):
                if not math.isfinite(_number(row.get(key))):
                    raise ValueError(f"{summary_path.name} 第 {line_number} 条记录的 {key} 不是有效数值。")
            if not _text(row.get("position")) or not _direction(row.get("direction")):
                raise ValueError(f"{summary_path.name} 第 {line_number} 条记录缺少位置或加载方向。")
            normalized: dict[str, object] = dict(row)
            device_id = _text(row.get("device_sn"), fallback_device)
            normalized.update({
                "device_id": device_id,
                "session_id": session_id,
                "source_file": summary_path.name,
                "position_norm": {"front": "前", "middle": "中", "rear": "后", "back": "后"}.get(
                    _text(row.get("position")).lower(), _text(row.get("position"))),
                "direction_norm": _direction(row.get("direction")),
                "quality_norm": _text(row.get("quality")),
            })
            for name in (
                "repeat", "step_index", "target_n", "reference_median_n",
                "upper_model_median_n", "temperature_c",
            ):
                normalized[name] = _number(row.get(name))
            all_rows.append(normalized)
            metadata_by_device.setdefault(device_id, {}).update(metadata)
    if not all_rows:
        names = "、".join(path.name for path in paths[:3])
        raise ValueError(f"找到 {len(paths)} 个 summary CSV，但都只有表头或内容为空，0 条采样记录。\n{names}\n"
                         "请换选有采样数据的文件或其目录；启动标定后未完成稳定采样也会生成只有表头的文件。")
    return all_rows, sources, metadata_by_device


def _mark_post_unload_zero(rows: list[dict[str, object]]) -> None:
    groups: dict[tuple[object, ...], list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        key = (row["device_id"], row["session_id"], row["repeat"], row["position_norm"])
        groups[key].append(row)
    for group in groups.values():
        has_loaded = False
        for row in sorted(group, key=lambda item: _number(item.get("step_index"), 0.0)):
            reference = _number(row.get("reference_median_n"))
            row["post_unload_zero"] = bool(has_loaded and reference == 0.0)
            if reference > 0.0:
                has_loaded = True


def build_metrics(
    rows: list[dict[str, object]],
    metadata_by_device: dict[str, dict[str, str]],
) -> dict[str, list[dict[str, object]]]:
    _mark_post_unload_zero(rows)
    for row in rows:
        reference = _number(row.get("reference_median_n"))
        output = _number(row.get("upper_model_median_n"))
        row["error_n"] = output - reference
        row["abs_error_n"] = abs(output - reference)
        row["loaded"] = reference > 0.0
        row["low_force"] = 0.0 < reference <= 2.0

    repeat_values: dict[str, list[float]] = defaultdict(list)
    repeat_groups: dict[tuple[object, ...], list[float]] = defaultdict(list)
    for row in rows:
        if row["loaded"]:
            key = (
                row["device_id"], row["position_norm"], row["direction_norm"],
                row["reference_median_n"],
            )
            repeat_groups[key].append(_number(row.get("upper_model_median_n")))
    for key, values in repeat_groups.items():
        value = _std(values)
        if value is not None:
            repeat_values[str(key[0])].append(value)

    hysteresis: dict[str, list[float]] = defaultdict(list)
    hyst_groups: dict[tuple[object, ...], list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        if row["loaded"]:
            key = (
                row["device_id"], row["session_id"], row["repeat"],
                row["position_norm"], row["reference_median_n"],
            )
            hyst_groups[key].append(row)
    for key, group in hyst_groups.items():
        loaded = [_number(row["upper_model_median_n"]) for row in group if row["direction_norm"] == "加载"]
        unloaded = [_number(row["upper_model_median_n"]) for row in group if row["direction_norm"] == "卸载"]
        if loaded and unloaded:
            hysteresis[str(key[0])].append(abs(float(median(loaded)) - float(median(unloaded))))

    position_spreads: dict[str, list[float]] = defaultdict(list)
    spread_groups: dict[tuple[object, ...], list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        if row["loaded"]:
            key = (
                row["device_id"], row["session_id"], row["repeat"],
                row["direction_norm"], row["reference_median_n"],
            )
            spread_groups[key].append(row)
    for key, group in spread_groups.items():
        by_position: dict[str, list[float]] = defaultdict(list)
        for row in group:
            by_position[str(row["position_norm"])].append(_number(row["upper_model_median_n"]))
        medians = [float(median(values)) for values in by_position.values() if values]
        if {"前", "中", "后"}.issubset(by_position):
            position_spreads[str(key[0])].append(max(medians) - min(medians))

    by_device: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        by_device[str(row["device_id"])].append(row)
    device_metrics: list[dict[str, object]] = []
    for device_id, group in sorted(by_device.items()):
        loaded = [row for row in group if row["loaded"]]
        low = [row for row in group if row["low_force"]]
        zero = [row for row in group if row.get("post_unload_zero")]
        abs_errors = [_number(row["abs_error_n"]) for row in loaded]
        zero_values = [abs(_number(row["upper_model_median_n"])) for row in zero]
        meta = metadata_by_device.get(device_id, {})
        device_metrics.append({
            "device_id": device_id,
            "silicone_batch": meta.get("silicone_batch", ""),
            "magnet_assembly": meta.get("magnet_assembly", ""),
            "operator": meta.get("operator", ""),
            "session_count": len({row["session_id"] for row in group}),
            "repeat_count": len({row["repeat"] for row in group}),
            "position_count": len({row["position_norm"] for row in group}),
            "loaded_steps": len(loaded),
            "overall_mae_n": _mean(abs_errors),
            "low_force_mae_n": _mean([_number(row["abs_error_n"]) for row in low]),
            "within_1n_rate": _mean([1.0 if _number(row["abs_error_n"]) <= 1.0 else 0.0 for row in loaded]),
            "p95_abs_error_n": _percentile(abs_errors, 95),
            "max_abs_error_n": max(abs_errors) if abs_errors else None,
            "repeatability_p95_n": _percentile(repeat_values[device_id], 95),
            "hysteresis_p95_n": _percentile(hysteresis[device_id], 95),
            "residual_median_n": _median(zero_values),
            "residual_p95_n": _percentile(zero_values, 95),
            "residual_max_n": max(zero_values) if zero_values else None,
            "position_spread_p95_n": _percentile(position_spreads[device_id], 95),
            "bad_quality_steps": sum(1 for row in group if str(row["quality_norm"]).upper() != "OK"),
        })

    position_metrics: list[dict[str, object]] = []
    position_groups: dict[tuple[str, str], list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        if row["loaded"]:
            position_groups[(str(row["device_id"]), str(row["position_norm"]))].append(row)
    for (device_id, position), group in position_groups.items():
        low = [row for row in group if row["low_force"]]
        errors = [_number(row["error_n"]) for row in group]
        position_metrics.append({
            "device_id": device_id,
            "position": position,
            "steps": len(group),
            "mae_n": _mean(abs(value) for value in errors),
            "low_force_mae_n": _mean(_number(row["abs_error_n"]) for row in low),
            "bias_n": _mean(errors),
            "within_1n_rate": _mean(1.0 if abs(value) <= 1.0 else 0.0 for value in errors),
            "p95_abs_error_n": _percentile((abs(value) for value in errors), 95),
        })
    position_metrics.sort(key=lambda item: (item["device_id"], POSITION_ORDER.get(str(item["position"]), 99)))

    force_points: list[dict[str, object]] = []
    point_groups: dict[tuple[object, ...], list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        if row["loaded"]:
            key = (
                row["device_id"], row["position_norm"], row["direction_norm"],
                row["reference_median_n"],
            )
            point_groups[key].append(row)
    for key, group in point_groups.items():
        outputs = [_number(row["upper_model_median_n"]) for row in group]
        errors = [_number(row["error_n"]) for row in group]
        force_points.append({
            "device_id": key[0], "position": key[1], "direction": key[2],
            "reference_n": key[3], "samples": len(group),
            "output_mean_n": _mean(outputs), "output_std_n": _std(outputs),
            "bias_n": _mean(errors), "mae_n": _mean(abs(value) for value in errors),
        })
    force_points.sort(key=lambda item: (
        item["device_id"], POSITION_ORDER.get(str(item["position"]), 99),
        item["direction"], item["reference_n"],
    ))

    consistency: list[dict[str, object]] = []
    consistency_groups: dict[tuple[object, ...], dict[str, list[float]]] = defaultdict(lambda: defaultdict(list))
    for row in rows:
        if row["loaded"]:
            key = (row["position_norm"], row["direction_norm"], row["reference_median_n"])
            consistency_groups[key][str(row["device_id"])].append(_number(row["upper_model_median_n"]))
    for key, device_values in consistency_groups.items():
        medians = [float(median(values)) for values in device_values.values() if values]
        if len(medians) < 2:
            continue
        device_std = float(stdev(medians))
        reference = float(key[2])
        consistency.append({
            "position": key[0], "direction": key[1], "reference_n": reference,
            "device_count": len(medians), "device_mean_n": _mean(medians),
            "device_std_n": device_std,
            "cv_percent_of_reference": device_std / max(abs(reference), 1e-9) * 100.0,
            "device_spread_n": max(medians) - min(medians),
        })
    consistency.sort(key=lambda item: (
        POSITION_ORDER.get(str(item["position"]), 99), item["direction"], item["reference_n"],
    ))
    return {
        "device_metrics": device_metrics,
        "position_metrics": position_metrics,
        "force_points": force_points,
        "device_consistency": consistency,
        "raw_steps": rows,
    }


def _sheet_title(sheet, title: str, subtitle: str, columns: int) -> int:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    sheet["A1"] = title
    sheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    sheet["A1"].font = Font(color=WHITE, bold=True, size=16)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    sheet["A2"] = subtitle
    sheet["A2"].fill = PatternFill("solid", fgColor=PALE)
    sheet["A2"].font = Font(color=BLUE, italic=True, size=10)
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[2].height = 34
    return 4


def _write_table(sheet, start_row: int, headers: list[str], rows: list[list[object]], name: str) -> tuple[int, int]:
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(start_row, column, header)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_index, values in enumerate(rows, start=start_row + 1):
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column, value)
            cell.fill = PatternFill("solid", fgColor=WHITE if row_index % 2 else PALE)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color=GRID))
    end_row = start_row + max(len(rows), 1)
    if rows:
        ref = f"A{start_row}:{get_column_letter(len(headers))}{end_row}"
        table = Table(displayName=name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        sheet.add_table(table)
    sheet.freeze_panes = f"A{start_row + 1}"
    sheet.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{end_row}"
    return start_row, end_row


def _fit_columns(sheet, max_width: int = 36) -> None:
    for column_cells in sheet.columns:
        letter = get_column_letter(column_cells[0].column)
        width = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=8)
        sheet.column_dimensions[letter].width = min(max(width + 2, 10), max_width)


def _format_numbers(sheet, start_row: int, end_row: int, percent_columns: set[int] | None = None) -> None:
    percent_columns = percent_columns or set()
    for row in sheet.iter_rows(min_row=start_row, max_row=end_row):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.0%" if cell.column in percent_columns else "0.000"


def write_workbook(
    output_path: str | Path,
    metrics: dict[str, list[dict[str, object]]],
    sources: list[dict[str, object]],
    input_dir: str | Path,
    progress: Callable[[int, str], None] | None = None,
) -> Path:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    summary = workbook.active
    summary.title = "综合评价"

    parameters = workbook.create_sheet("判定参数")
    row = _sheet_title(parameters, "H3 一致性评价参数", "黄色单元格可直接修改；设备判定、得分和排名将随公式自动更新。", 5)
    parameter_rows = [
        ["overall_mae_n", "总体 MAE 上限", 1.0, "N", "越小越好"],
        ["low_force_mae_n", "≤2 N MAE 上限", 0.5, "N", "重点考核低力"],
        ["p95_abs_error_n", "绝对误差 P95 上限", 2.0, "N", "越小越好"],
        ["repeatability_p95_n", "重复性标准差 P95 上限", 0.5, "N", "同条件跨轮"],
        ["hysteresis_p95_n", "迟滞 P95 上限", 1.0, "N", "加载/卸载差"],
        ["residual_median_n", "卸载残余中位数上限", 0.1, "N", "回零重点"],
        ["residual_max_n", "卸载残余最大值上限", 1.5, "N", "快速回落要求"],
        ["position_spread_p95_n", "前中后差异 P95 上限", 2.0, "N", "位置一致性"],
        ["within_1n_rate", "误差≤1 N 比例下限", 0.8, "比例", "越大越好"],
        ["bad_quality_steps", "异常质量步数上限", 0, "步", "应为零"],
    ]
    p_start, p_end = _write_table(parameters, row, ["参数键", "判定项", "阈值", "单位", "说明"], parameter_rows, "Thresholds")
    for cell in parameters.iter_rows(min_row=p_start + 1, max_row=p_end, min_col=3, max_col=3):
        cell[0].fill = PatternFill("solid", fgColor=YELLOW)
        cell[0].font = Font(color="7F6000", bold=True)
    _fit_columns(parameters)

    device_sheet = workbook.create_sheet("设备指标")
    d_headers = [
        "设备", "硅胶批次", "磁铁组件", "操作员", "会话数", "轮数", "位置数", "有效力点",
        "总体MAE(N)", "≤2N MAE(N)", "误差≤1N比例", "绝对误差P95(N)", "最大误差(N)",
        "重复性P95(N)", "迟滞P95(N)", "卸载残余中位数(N)", "卸载残余P95(N)",
        "卸载残余最大值(N)", "前中后差异P95(N)", "异常质量步", "综合判定", "选优得分", "排名",
    ]
    d_row = _sheet_title(device_sheet, "设备指标与选优", "每个 device_sn 作为一台设备；阈值来自“判定参数”。只有一台设备时可评价自身，不能证明设备间一致性。", len(d_headers))
    d_rows: list[list[object]] = []
    for index, item in enumerate(metrics["device_metrics"], start=d_row + 1):
        values = [
            item["device_id"], item["silicone_batch"], item["magnet_assembly"], item["operator"],
            item["session_count"], item["repeat_count"], item["position_count"], item["loaded_steps"],
            item["overall_mae_n"], item["low_force_mae_n"], item["within_1n_rate"],
            item["p95_abs_error_n"], item["max_abs_error_n"], item["repeatability_p95_n"],
            item["hysteresis_p95_n"], item["residual_median_n"], item["residual_p95_n"],
            item["residual_max_n"], item["position_spread_p95_n"], item["bad_quality_steps"],
        ]
        status = (
            f'=IF(COUNT(I{index}:L{index},N{index}:P{index},R{index}:T{index})<10,"数据不足",IF(AND(I{index}<=\'判定参数\'!$C$5,J{index}<=\'判定参数\'!$C$6,'
            f'K{index}>=\'判定参数\'!$C$13,L{index}<=\'判定参数\'!$C$7,'
            f'N{index}<=\'判定参数\'!$C$8,O{index}<=\'判定参数\'!$C$9,'
            f'P{index}<=\'判定参数\'!$C$10,R{index}<=\'判定参数\'!$C$11,'
            f'S{index}<=\'判定参数\'!$C$12,T{index}<=\'判定参数\'!$C$14),"合格","不合格"))'
        )
        score = (
            f'=ROUND(100-MIN(100,I{index}/\'判定参数\'!$C$5*15)'
            f'-MIN(100,J{index}/\'判定参数\'!$C$6*20)'
            f'-MIN(100,P{index}/\'判定参数\'!$C$10*20)'
            f'-MIN(100,R{index}/\'判定参数\'!$C$11*15)'
            f'-MIN(100,N{index}/\'判定参数\'!$C$8*10)'
            f'-MIN(100,O{index}/\'判定参数\'!$C$9*10)'
            f'-MIN(100,S{index}/\'判定参数\'!$C$12*10),1)'
        )
        score = f'=IF(U{index}="数据不足","",{score[1:]})'
        rank = f'=IF(ISNUMBER(V{index}),RANK(V{index},$V${d_row + 1}:$V${d_row + len(metrics["device_metrics"])},0),"")'
        d_rows.append([*values, status, score, rank])
    d_start, d_end = _write_table(device_sheet, d_row, d_headers, d_rows, "DeviceMetrics")
    _format_numbers(device_sheet, d_start + 1, d_end, {11})
    if d_rows:
        device_sheet.conditional_formatting.add(
            f"U{d_start + 1}:U{d_end}", CellIsRule(operator="equal", formula=['"数据不足"'], fill=PatternFill("solid", fgColor=YELLOW))
        )
        device_sheet.conditional_formatting.add(
            f"U{d_start + 1}:U{d_end}", CellIsRule(operator="equal", formula=['"合格"'], fill=PatternFill("solid", fgColor=GREEN))
        )
        device_sheet.conditional_formatting.add(
            f"U{d_start + 1}:U{d_end}", CellIsRule(operator="equal", formula=['"不合格"'], fill=PatternFill("solid", fgColor=RED))
        )
    _fit_columns(device_sheet, 24)

    row = _sheet_title(summary, "H3 设备一致性评价", "导入标定 summary/frames 后自动生成。结论基于当前试行阈值，正式供样前请由项目负责人确认阈值。", 8)
    summary["A4"] = "生成时间"
    summary["B4"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    summary["A5"] = "数据目录"
    summary.merge_cells("B5:H5")
    summary["B5"] = str(Path(input_dir).resolve())
    summary["A6"] = "识别设备数"
    summary["B6"] = len(metrics["device_metrics"])
    summary["D6"] = "有效会话数"
    summary["E6"] = sum(1 for source in sources if source["summary_rows"])
    summary["G6"] = "空文件数"
    summary["H6"] = sum(1 for source in sources if not source["summary_rows"])
    summary["A8"] = "总体结论"
    if len(metrics["device_metrics"]) == 1:
        summary["B8"] = '=IF(\'设备指标\'!U5="数据不足","数据不足，不能给出完整合格判定；尚缺第二台设备，不能评价设备间一致性",IF(\'设备指标\'!U5="合格","单机指标合格；尚缺第二台设备，不能评价设备间一致性","单机指标不合格；请查看设备指标与位置/力点明细"))'
    else:
        summary["B8"] = '=IF(COUNTIF(\'设备指标\'!U:U,"数据不足")>0,"存在数据不足的设备，不能给出完整判定",IF(COUNTIF(\'设备指标\'!U:U,"不合格")=0,"所有设备均符合当前试行阈值","存在不合格设备，请查看设备指标和选优排名"))'
    summary.merge_cells("B8:H8")
    summary["B8"].fill = PatternFill("solid", fgColor=YELLOW)
    summary["B8"].font = Font(color="7F6000", bold=True, size=12)
    summary["B8"].alignment = Alignment(wrap_text=True)
    summary.row_dimensions[8].height = 48
    summary["A10"] = "使用步骤"
    summary["A10"].font = Font(color=BLUE, bold=True, size=12)
    instructions = [
        "1. 每台设备使用唯一 device_sn，按同一采集方案完成前/中/后多轮标定。",
        "2. 在上位机点击“一致性评价 Excel”，选择包含 CSV 的总目录。",
        "3. 打开本报告，先看综合评价，再看设备指标、位置与力点、设备间一致性。",
        "4. 黄色阈值可修改；缺失指标显示数据不足，不计分、不排名。空白不是零。",
    ]
    for offset, instruction in enumerate(instructions, start=11):
        summary.merge_cells(start_row=offset, start_column=1, end_row=offset, end_column=8)
        summary.cell(offset, 1, instruction)
        summary.cell(offset, 1).alignment = Alignment(wrap_text=True)
    if metrics["device_metrics"]:
        chart = BarChart()
        chart.title = "设备关键误差比较"
        chart.y_axis.title = "N"
        chart.x_axis.title = "设备"
        data = Reference(device_sheet, min_col=9, max_col=10, min_row=d_start, max_row=d_end)
        categories = Reference(device_sheet, min_col=1, min_row=d_start + 1, max_row=d_end)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        chart.height = 8
        chart.width = 15
        summary.add_chart(chart, "A17")
    for column, width in zip(range(1, 9), [18, 22, 18, 18, 18, 18, 18, 18]):
        summary.column_dimensions[get_column_letter(column)].width = width
    summary.freeze_panes = "A4"

    position_sheet = workbook.create_sheet("位置与力点")
    pos_headers = ["设备", "位置", "有效步", "MAE(N)", "≤2N MAE(N)", "偏差(N)", "误差≤1N比例", "绝对误差P95(N)"]
    pos_row = _sheet_title(position_sheet, "前 / 中 / 后位置评价", "用于判断同一硬件在三个按压区域的精度差异。", len(pos_headers))
    pos_rows = [[item[key] for key in ("device_id", "position", "steps", "mae_n", "low_force_mae_n", "bias_n", "within_1n_rate", "p95_abs_error_n")] for item in metrics["position_metrics"]]
    pos_start, pos_end = _write_table(position_sheet, pos_row, pos_headers, pos_rows, "PositionMetrics")
    _format_numbers(position_sheet, pos_start + 1, pos_end, {7})
    force_row = pos_end + 3
    force_headers = ["设备", "位置", "方向", "参考力(N)", "样本数", "输出均值(N)", "输出标准差(N)", "偏差(N)", "MAE(N)"]
    force_rows = [[item[key] for key in ("device_id", "position", "direction", "reference_n", "samples", "output_mean_n", "output_std_n", "bias_n", "mae_n")] for item in metrics["force_points"]]
    force_start, force_end = _write_table(position_sheet, force_row, force_headers, force_rows, "ForcePoints")
    _format_numbers(position_sheet, force_start + 1, force_end)
    _fit_columns(position_sheet, 22)

    consistency_sheet = workbook.create_sheet("设备间一致性")
    c_headers = ["位置", "方向", "参考力(N)", "设备数", "设备均值(N)", "设备标准差(N)", "CV/参考力(%)", "设备极差(N)"]
    c_row = _sheet_title(consistency_sheet, "设备间一致性", "只有相同位置、方向和参考力下至少两台不同 device_sn 才会产生数据。", len(c_headers))
    c_rows = [[item[key] for key in ("position", "direction", "reference_n", "device_count", "device_mean_n", "device_std_n", "cv_percent_of_reference", "device_spread_n")] for item in metrics["device_consistency"]]
    c_start, c_end = _write_table(consistency_sheet, c_row, c_headers, c_rows, "DeviceConsistency")
    if not c_rows:
        consistency_sheet["A6"] = "当前仅有一台设备或无可对齐力点；请增加不同 device_sn 的同方案数据。"
        consistency_sheet.merge_cells("A6:H6")
        consistency_sheet["A6"].fill = PatternFill("solid", fgColor=YELLOW)
    _format_numbers(consistency_sheet, c_start + 1, c_end)
    _fit_columns(consistency_sheet, 24)

    raw_sheet = workbook.create_sheet("原始步汇总")
    raw_headers = ["设备", "会话", "完成时间", "轮次", "位置", "方向", "目标力(N)", "参考力(N)", "模型输出(N)", "误差(N)", "绝对误差(N)", "温度(℃)", "质量", "卸载后零点", "来源文件"]
    raw_row = _sheet_title(raw_sheet, "原始稳定步汇总", "保留每个稳定采样步骤，便于追溯异常点。", len(raw_headers))
    raw_rows = []
    for item in sorted(metrics["raw_steps"], key=lambda value: (str(value["device_id"]), str(value["session_id"]), _number(value.get("step_index"), 0.0))):
        raw_rows.append([
            item["device_id"], item["session_id"], item.get("completed_time_iso", ""), item["repeat"],
            item["position_norm"], item["direction_norm"], item["target_n"], item["reference_median_n"],
            item["upper_model_median_n"], item["error_n"], item["abs_error_n"], item["temperature_c"],
            item["quality_norm"], "是" if item.get("post_unload_zero") else "否", item["source_file"],
        ])
    raw_start, raw_end = _write_table(raw_sheet, raw_row, raw_headers, raw_rows, "RawSteps")
    _format_numbers(raw_sheet, raw_start + 1, raw_end)
    _fit_columns(raw_sheet, 30)

    source_sheet = workbook.create_sheet("来源说明")
    source_headers = ["会话", "summary 文件", "frames 文件", "summary 行", "frames 行", "旧表头已修复", "状态"]
    source_row = _sheet_title(source_sheet, "数据来源与导入状态", "原始 CSV 不会被修改；空文件会保留记录并自动跳过。", len(source_headers))
    source_rows = [[source[key] for key in ("session_id", "summary_file", "frames_file", "summary_rows", "frame_rows", "header_recovered", "status")] for source in sources]
    source_start, source_end = _write_table(source_sheet, source_row, source_headers, source_rows, "Sources")
    _format_numbers(source_sheet, source_start + 1, source_end)
    _fit_columns(source_sheet, 55)

    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.auto_filter.ref = sheet.auto_filter.ref
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    if progress:
        progress(85, "写入 Excel 工作簿")
    workbook.save(path)
    if progress:
        progress(100, f"报告已生成：{path.name}")
    return path


def generate_consistency_report(
    input_dir: str | Path,
    output_path: str | Path,
    progress: Callable[[int, str], None] | None = None,
    mode: str = "calibration",
) -> dict[str, object]:
    if mode in ("folders", "folders_md"):
        from folder_report import generate_folder_reports
        return generate_folder_reports(input_dir, output_path, progress, output_format='md' if mode == 'folders_md' else 'xlsx')
    if mode == "recordings":
        from recording_report import generate_recording_report
        return generate_recording_report(input_dir, output_path, progress)
    if mode != "calibration":
        raise ValueError(f"未知报告类型：{mode}")
    rows, sources, metadata = load_sessions(input_dir, progress)
    if progress:
        progress(42, "计算精度、低力、重复性、迟滞和卸载残余")
    metrics = build_metrics(rows, metadata)
    if progress:
        progress(70, "生成设备一致性和选优排名")
    path = write_workbook(output_path, metrics, sources, input_dir, progress)
    return {
        "path": str(path),
        "device_count": len(metrics["device_metrics"]),
        "session_count": sum(1 for source in sources if source["summary_rows"]),
        "skipped_empty": sum(1 for source in sources if not source["summary_rows"]),
        "recovered_headers": sum(1 for source in sources if source["header_recovered"]),
    }


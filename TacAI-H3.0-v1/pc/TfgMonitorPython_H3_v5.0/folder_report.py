"""One folder, one Sheet1 workbook; one raw CSV, one row of region means.

The formulas and row positions follow the user's 本次测试数据结果分析表.
H3 S1-S8 allocations retain their names and are never relabelled as 12 sensors.
"""
from __future__ import annotations

from collections import Counter
import csv
import hashlib
import io
import math
from pathlib import Path
import re
from statistics import mean
import tempfile
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from recording_report import (FLAG_LABELS, acquisition_diagnostics, audit_effective_baseline,
                              continuous_protocol_reason, csv_paths, invalid_paired_capture)
from baseline_quality import audit_saved_baselines

H3_COLUMNS = [f"S{i}_resultant_allocation_estimated_n" for i in range(1, 9)]
NO7_COLUMNS = [f"P{i:02}_estimated_fz_n" for i in range(12)]
FS = 30.0

# Compare only metadata, never the per-frame force or correction amounts.
SHARED_CONFIG_COLUMNS = (
    "selected_value_mode", "scalar_model_status",
    "spatial_allocation_method", "spatial_allocation_exponent",
)
NUMERIC_CONFIG_COLUMNS = {
    "spatial_allocation_exponent",
    *(f"S{i}_spatial_noise_floor" for i in range(1, 9)),
}


class ConfigurationConflict(ValueError):
    """A mixed configuration invalidates the group, not just one trial."""


def sample_configuration(path, header, samples):
    """Audit named metadata on valid paired rows; old NO7 events stay intact."""
    paired = [row for row in samples
              if row.get("record_type", "PAIRED").strip().upper() == "PAIRED"]
    if not paired:
        return {}
    fields = [field for field in header if field in SHARED_CONFIG_COLUMNS
              or field in NUMERIC_CONFIG_COLUMNS
              or field == "correction_signature"
              or field.endswith("_correction_signature")]
    result = {}
    for field in fields:
        for row in paired:
            raw = row[field].strip()
            value = raw
            if raw and field in NUMERIC_CONFIG_COLUMNS:
                value = finite_number(raw)
                if value is None:
                    raise ConfigurationConflict(
                        f"{path.name} 第 {row['_source_line']} 行：配置 {field}={raw!r} 不是有限数值；不能混合评价。")
            current = (value, raw, row["_source_line"])
            previous = result.setdefault(field, current)
            if previous[0] != value:
                raise ConfigurationConflict(
                    f"{path.name} 同一文件配置变化：{field}，"
                    f"第 {previous[2]} 行 {previous[1]!r} → 第 {current[2]} 行 {raw!r}；"
                    "请按相同配置分别采集和评价，本文件夹未生成结果。")
    return result


def compare_folder_configuration(path, current, shared):
    # Per-trial tare/noise floors may legitimately change after unloading.
    # Absent legacy columns cannot establish a conflicting configuration.
    fields = set(SHARED_CONFIG_COLUMNS)
    fields.update(field for field in current
                  if field == "correction_signature" or field.endswith("_correction_signature"))
    for field in sorted(fields):
        if field not in current:
            continue
        value, raw, line = current[field]
        previous = shared.setdefault(field, (value, raw, line, path.name))
        if previous[0] != value:
            raise ConfigurationConflict(
                f"同一文件夹配置不一致：{field}，"
                f"{previous[3]} 第 {previous[2]} 行 {previous[1]!r}，"
                f"{path.name} 第 {line} 行 {raw!r}；"
                "请按相同模型、模式、分配方法和指数分别评价，本文件夹未生成结果。")


def read_region_csv(path):
    """Read named columns; Excel-added unnamed columns are not region data."""
    data = path.read_bytes()
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("gb18030")
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return [], [], data
    header = [value.strip() for value in rows[0]]
    named = [value for value in header if value]
    if len(set(named)) != len(named):
        raise ValueError("CSV 存在重复的非空列名，请核对原始表头。")
    records = []
    last_named = max((i for i, value in enumerate(header) if value), default=-1)
    for line, row in enumerate(rows[1:], 2):
        if not any(value.strip() for value in row):
            continue
        if len(row) > len(header) or len(row) <= last_named:
            raise ValueError(f"第 {line} 行字段数与表头不符，无法可靠对应区域列。")
        record = {name: row[i] for i, name in enumerate(header) if name}
        record["_source_line"] = line
        records.append(record)
    return header, records, data


def finite_number(value):
    try:
        value = float(value)
        return value if math.isfinite(value) else None
    except (ValueError, TypeError):
        return None


def region_columns(header):
    if all(column in header for column in H3_COLUMNS):
        return [f"S{i}" for i in range(1, 9)], H3_COLUMNS, "H3 八通道估计分配值（N）"
    if all(column in header for column in NO7_COLUMNS):
        return [f"P{i:02}" for i in range(12)], NO7_COLUMNS, "NO7 / H2.2 十二显示区估计值（N）"
    columns = {}
    for column in header:
        match = re.fullmatch(r"P(\d{2})_Fz(?:\([^)]*\))?", column, re.I)
        if match and int(match[1]) < 12:
            index = int(match[1])
            if index in columns:
                raise ValueError(f"同一区域有多个 Fz 列：P{index:02}")
            columns[index] = column
    if len(columns) == 12:
        return [f"P{i:02}" for i in range(12)], [columns[i] for i in range(12)], "原始 CSV 的 P00–P11 Fz（N）"
    raise ValueError("缺少区域列：需要 S1–S8 分配值、P00–P11_estimated_fz_n 或完整的 P00–P11_Fz 列；summary 不属于单轮原始数据。")


def natural_key(path):
    return [int(part) if part.isdigit() else part.casefold() for part in re.split(r"(\d+)", path.name)]


def load_folder(folder):
    paths = sorted((p for p in Path(folder).iterdir() if p.is_file() and p.suffix.lower() == ".csv"), key=natural_key)
    if not paths:
        raise ValueError("文件夹中没有 CSV。")
    if len(paths) > 10:
        raise ValueError(f"本文件夹有 {len(paths)} 个 CSV；模板为 10 次，请把每组 10 次分别放到一个文件夹。")
    ordinals = [re.search(r"[_ -](\d{1,2})$", p.stem) for p in paths]
    numbers = [int(m[1]) if m else None for m in ordinals]
    use_numbers = all(n is not None and 1 <= n <= 10 for n in numbers) and len(set(numbers)) == len(numbers)
    entries, schema, seen_hashes, shared_config = [], None, {}, {}
    for index, path in enumerate(paths, 1):
        entry = {"round": numbers[index-1] if use_numbers else index, "path": str(path), "name": path.name, "status": "未导入",
                 "means": [], "counts": [], "rows": 0, "sample_rows": 0, "warnings": "", "error": "",
                 "selection": "", "footer_check": "未发现附加均值行", "footer_max_error": None,
                 "invalid_cycles": 0, "acquisition_events": 0, "diagnostics": ""}
        entries.append(entry)
        try:
            header, rows, data = read_region_csv(path)
            diagnostics = acquisition_diagnostics(rows)
            entry.update(rows=len(rows), invalid_cycles=diagnostics["invalid_cycles"],
                         acquisition_events=diagnostics["event_count"], diagnostics=diagnostics["text"],
                         warnings=diagnostics["text"], sha256=hashlib.sha256(data).hexdigest())
            labels, columns, meaning = region_columns(header)
            if not rows:
                raise ValueError("只有表头，没有采样数据。")
            is_no7 = columns == NO7_COLUMNS
            allowed = {"RAW", "FORCE", "PAIRED"} if is_no7 else {"PAIRED"}
            samples = [row for row in rows if row.get("record_type", "PAIRED").strip().upper() in allowed
                       and not invalid_paired_capture(row)]
            if not samples:
                raise ValueError("没有 RAW / FORCE / PAIRED 采样行。" if is_no7 else "没有 PAIRED 采样行。")
            continuous_reason = continuous_protocol_reason(samples)
            if continuous_reason:
                raise ConfigurationConflict(f"{path.name}：{continuous_reason} 本文件夹未生成结果。")
            try:
                working_zero_notice = audit_effective_baseline(samples)
            except ValueError as error:
                raise ConfigurationConflict(f"{path.name}：{error} 本文件夹未生成结果。") from error
            configuration = sample_configuration(path, header, samples)
            compare_folder_configuration(path, configuration, shared_config)
            tare_epochs = {row.get("host_tare_epoch", "").strip() for row in samples} - {""}
            if len(tare_epochs) > 1:
                raise ValueError("同一文件采样跨多个上位机置零基线（host_tare_epoch）；请按基线分别导出，不能混合评价。")
            types = Counter(row.get("record_type", "数值记录").strip().upper() for row in samples)
            selection = "＋".join(f"{kind} {count} 行" for kind, count in sorted(types.items()))
            if is_no7:
                selection += "；沿用 NO7：两类事件行均参与平均，末尾手工均值行不参与"
            if diagnostics["event_count"]:
                selection += f"；存在 {diagnostics['event_count']} 条采集异常事件，仅统计有效记录"
            warnings = [diagnostics["text"]] if diagnostics["text"] else []
            if working_zero_notice:warnings.append(working_zero_notice)
            baseline_audit=audit_saved_baselines(samples)
            entry['baseline_audit']=baseline_audit
            if baseline_audit['text']:warnings.append(baseline_audit['text'])
            counts, means = [], []
            for label, column in zip(labels, columns):
                values = []
                for row in samples:
                    value = finite_number(row[column])
                    if value is not None:
                        values.append(value)
                if not values:
                    raise ValueError(f"{label} 没有可计算的数值。")
                means.append(mean(values))
                counts.append(len(values))
                if len(values) < len(samples):
                    warnings.append(f"{label} 有 {len(samples)-len(values)} 个空值/非有限数值未参与该列平均")
            current_schema = (labels, columns, meaning)
            if schema is not None and schema != current_schema:
                raise ValueError("同一文件夹内的区域列不一致，请按相同格式分开存放。")
            schema = current_schema
            if "record_type" in header:
                # The user's NO7 CSVs have a blank row then a saved mean row.
                # Recompute from events and use the saved row only as a check.
                added = [row for row in rows if not row["record_type"].strip()
                         and any(finite_number(row[column]) is not None for column in columns)]
                if added:
                    saved = [finite_number(added[-1][column]) for column in columns]
                    errors = [abs(value-means[i]) for i, value in enumerate(saved) if value is not None]
                    maximum = max(errors)
                    count = len(errors)
                    verdict = "一致（小数舍入范围内）" if count == len(columns) and maximum <= 1e-8 else "有差异或缺值，请核对"
                    entry["footer_max_error"] = maximum
                    entry["footer_check"] = f"{verdict}；{count}/{len(columns)} 区，最大差 {maximum:.9g} N；{len(added)} 行附加数值未计入采样"
            flags = Counter(flag.strip() for row in samples for flag in row.get("quality_flags", "").split("|") if flag.strip() and flag.strip() != "OK")
            for flag, count in sorted(flags.items()):
                warnings.append(f"{FLAG_LABELS.get(flag, flag)}：{count} 行（保留数值）")
            states = Counter(row.get("frame_state", "").strip() for row in samples)
            for state, count in states.items():
                if state and state != "PAIRED":
                    warnings.append(f"帧状态 {state}：{count} 行（保留数值）")
            digest = hashlib.sha256(data).hexdigest()
            if digest in seen_hashes:
                warnings.append(f"内容与 {seen_hashes[digest]} 完全相同，请核对是否重复文件")
            else:
                seen_hashes[digest] = path.name
            entry.update(status="已导入", means=means, counts=counts, rows=len(rows), sample_rows=len(samples), warnings="；".join(warnings), selection=selection,
                         samples=samples, sha256=digest,
                         source_columns="；".join(f"{label}={get_column_letter(header.index(column)+1)}列" for label, column in zip(labels, columns)))
        except ConfigurationConflict:
            # Do not omit a conflicting trial and silently report the rest.
            raise
        except (ValueError, OSError, UnicodeError) as exc:
            entry["error"] = str(exc)
    if schema is None:
        raise ValueError("本文件夹没有可计算的区域数据。" + "；".join(f"{e['name']}：{e['error']}" + ("；" + e["diagnostics"] if e["diagnostics"] else "") for e in entries))
    return schema, entries


def _save_with_caches(book, output, cache):
    """Cache the known Sheet1 formula results for viewers without a calc engine.

    Formulas remain editable. Excel recalculates on open; caches are calculated
    from the same unrounded input means, and independently verified in tests.
    """
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with tempfile.TemporaryDirectory(dir=output.parent) as temporary:
        raw = Path(temporary) / "raw.xlsx"
        cached = Path(temporary) / "cached.xlsx"
        book.save(raw)
        with ZipFile(raw) as source, ZipFile(cached, "w") as target:
            for item in source.infolist():
                content = source.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    root = ET.fromstring(content)
                    for cell in root.iter(f"{{{namespace}}}c"):
                        if cell.find(f"{{{namespace}}}f") is None:
                            continue
                        value = cache.get(cell.attrib["r"])
                        node = cell.find(f"{{{namespace}}}v")
                        if node is None:
                            node = ET.SubElement(cell, f"{{{namespace}}}v")
                        if value is None:
                            cell.set("t", "str")
                            node.text = ""
                        else:
                            cell.attrib.pop("t", None)
                            node.text = repr(float(value))
                    content = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                target.writestr(item, content)
        cached.replace(output)


def write_folder_report(folder, output, schema, entries):
    labels, columns, meaning = schema
    book = Workbook()
    book.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True)
    sheet = book.active
    sheet.title = "Sheet1"
    last = get_column_letter(len(labels) + 1)
    sheet["A1"] = "FS=30N"
    sheet["B1"] = Path(folder).name
    sheet.merge_cells(f"C1:{last}1")
    sheet["C1"] = meaning
    sheet["A2"] = "测试轮次"
    for row, text in [(17, "差值表"), (29, "FS%"), (41, "FS%MAX"), (43, "FS%SUM_MAX"), (45, "FS%mean"), (47, "FS%SUM_mean")]:
        sheet.cell(row, 1, text)
    for i in range(10):
        for offset in (3, 18, 30):
            sheet.cell(offset + i, 1, i + 1)
    for row in (42, 44, 46, 48):
        sheet.cell(row, 1, 1)
    by_round = {entry["round"]: entry for entry in entries}
    # Make every mean inspectable in Excel without changing sample selection.
    detail = book.create_sheet("输入明细")
    trace = book.create_sheet("逐帧区域值")
    trace.append(["测试轮次", "CSV原始行号", "记录类型", "采样时间", "device_ms", *columns, "质量标记", "原始文件"])
    trace_ranges = {}
    for entry in sorted(entries, key=lambda item: item["round"]):
        if entry["status"] != "已导入":
            continue
        start = trace.max_row + 1
        for sample in entry["samples"]:
            trace.append([entry["round"], sample["_source_line"], sample.get("record_type", "PAIRED"),
                          sample.get("local_time_iso", ""), sample.get("device_ms", ""),
                          *[finite_number(sample[column]) for column in columns],
                          sample.get("quality_flags", ""), entry["name"]])
        trace_ranges[entry["round"]] = (start, trace.max_row)
    trace.freeze_panes = "F2"
    trace.auto_filter.ref = trace.dimensions
    for column, width in zip("ABCDE", (12, 16, 14, 27, 16)):
        trace.column_dimensions[column].width = width
    for index in range(len(labels)):
        trace.column_dimensions[get_column_letter(index+6)].width = 28
    trace.column_dimensions[get_column_letter(len(labels)+6)].width = 55
    trace.column_dimensions[get_column_letter(len(labels)+7)].width = 65
    trace.row_dimensions[1].height = 48
    for row in trace:
        for cell in row:
            cell.font = Font(name="宋体", size=11, bold=cell.row == 1)
            cell.alignment = Alignment(vertical="center", wrap_text=cell.row == 1)
            if cell.row > 1 and 6 <= cell.column < 6 + len(labels):
                cell.number_format = "0.000000000"
    cache, maxima, averages = {}, [], []
    for index, label in enumerate(labels):
        col = get_column_letter(index + 2)
        for row, suffix in [(2, "mean"), (14, "allmean"), (17, "dev")]:
            metric = "估计值" if columns == H3_COLUMNS else "Fz"
            sheet.cell(row, index + 2, f"{label}_{metric}({suffix})")
        if "estimated" in columns[index]:
            sheet.cell(2, index + 2).comment = Comment(f"来源列：{columns[index]}。这是合力的区域估计分配值，并非独立标定的区域力传感器读数。", "TacAI")
        values = []
        for i in range(10):
            entry = by_round.get(i + 1)
            if entry and entry["status"] == "已导入":
                value = entry["means"][index]
                first, final = trace_ranges[i+1]
                trace_col = get_column_letter(index+6)
                sheet.cell(i + 3, index + 2, f"=AVERAGE('逐帧区域值'!{trace_col}{first}:{trace_col}{final})")
                cache[f"{col}{i+3}"] = value
                sheet.cell(i + 3, index + 2).comment = Comment(f"文件：{entry['name']}\n源列：{columns[index]}\n参与平均：{entry['counts'][index]} 行；算术平均，未截取稳定区间。\n{entry['selection']}" + ("\n" + entry["diagnostics"] if entry["diagnostics"] else ""), "TacAI")
                values.append(value)
        overall = mean(values) if values else None
        sheet[f"{col}15"] = f'=IF(COUNT({col}3:{col}12)=0,"",AVERAGE({col}3:{col}12))'
        cache[f"{col}15"] = overall
        deviations = []
        for i in range(10):
            sheet[f"{col}{18+i}"] = f'=IF({col}{3+i}="","",ABS({col}{3+i}-{col}$15))'
            sheet[f"{col}{30+i}"] = f'=IF({col}{18+i}="","",{col}{18+i}/30)'
            value = cache.get(f"{col}{3+i}")
            deviation = abs(value - overall) if value is not None and overall is not None else None
            cache[f"{col}{18+i}"] = deviation
            cache[f"{col}{30+i}"] = deviation / FS if deviation is not None else None
            if deviation is not None:
                deviations.append(deviation / FS)
        sheet[f"{col}42"] = f'=IF(COUNT({col}30:{col}39)=0,"",MAX({col}30:{col}39))'
        sheet[f"{col}46"] = f'=IF(COUNT({col}30:{col}39)=0,"",AVERAGE({col}30:{col}39))'
        cache[f"{col}42"] = max(deviations) if deviations else None
        cache[f"{col}46"] = mean(deviations) if deviations else None
        if deviations:
            maxima.append(max(deviations))
            averages.append(mean(deviations))
    sheet["B44"] = f'=IF(COUNT(B42:{last}42)=0,"",SUM(B42:{last}42))'
    sheet["B48"] = f'=IF(COUNT(B46:{last}46)=0,"",SUM(B46:{last}46))'
    cache["B44"] = sum(maxima) if maxima else None
    cache["B48"] = sum(averages) if averages else None
    imported = sum(e["status"] == "已导入" for e in entries)
    notes = [f"输入：{Path(folder)}", f"{len(entries)} 个 CSV，已填 {imported}/10 行；一文件一轮，逐区域取全部有效采样数值的算术平均。",
             "区域数与原始 CSV 一致；空值、非数值不当作 0，详情见“输入明细”。",
             "差值=ABS(单轮区域均值−该区各轮均值)；FS%=差值÷30 N；不作合格判定。",
             "FS%SUM_MAX=各区 FS%MAX 相加；FS%SUM_mean=各区 FS%mean 相加（沿用原表）。",
             "每区参考指标：第42行最大偏差≤2%FS，第46行平均偏差≤0.5%FS；第44/48行总和不用于单区判定。",
             "双击单轮均值可查看 AVERAGE 来源；“逐帧区域值”保留原 CSV 行号、列名与质量提示。",
             "H3 的 S1–S8 是总力的空间分配估计；dBz 是磁信号，不是 Fz。区域力尚未分别独立标定。" if columns == H3_COLUMNS else
             "NO7 的 P00–P11 为旧程序的十二显示区估计值；保留原列，不从当前 H3 八通道扩造十二区。"]
    affected = [str(entry["round"]) for entry in entries if entry["acquisition_events"]]
    if affected:
        notes.append(f"第 {', '.join(affected)} 轮含采集异常：共 {sum(e['acquisition_events'] for e in entries)} 条事件，其中 {sum(e['invalid_cycles'] for e in entries)} 个明确无效采集周期。事件不补零或沿用旧值；结果仅覆盖有效记录，不能据此确认完整复测。诊断与 CSV 行号见输入明细。")
    for row, note in enumerate(notes, 50):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(labels) + 1)
        sheet.cell(row, 1, note)
        sheet.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[row].height = 48 if row >= 58 else (30 if row == 50 else 24)
    final_note_row = 49 + len(notes)
    for row in sheet.iter_rows(min_row=1, max_row=final_note_row, max_col=len(labels) + 1):
        for cell in row:
            cell.font = Font(name="宋体", size=11)
            if cell.row < 50:
                cell.alignment = Alignment(horizontal="left" if cell.row in (1, 2, 14, 17) or cell.column == 1 else "right", vertical="center")
            if cell.column > 1 and cell.row in [*range(3, 13), 15, *range(18, 28)]:
                cell.number_format = "0.000000000"
            if cell.column > 1 and cell.row in [*range(30, 40), 42, 44, 46, 48]:
                cell.number_format = "0.0000%"
    sheet.column_dimensions["A"].width = 19
    for index in range(len(labels)):
        sheet.column_dimensions[get_column_letter(index + 2)].width = 20
    for row in (2, 14, 17):
        sheet.row_dimensions[row].height = 24
    sheet.row_dimensions[1].height = 26
    sheet.freeze_panes = "B3"
    sheet.sheet_view.zoomScale = 85
    sheet.print_options.horizontalCentered = True
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_area = f"A1:{last}{final_note_row}"
    detail.append(["测试轮次", "原始文件", "状态", "采样行数", "各区参与平均行数", "数据提示／未导入原因", "完整路径", "记录类型／统计口径", "CSV 末尾均值核对", "原文件 SHA256", "原CSV列位置（完整列名见逐帧页）"])
    for i in range(1, 11):
        e = by_round.get(i)
        if e:
            message = "；".join(part for part in (e["error"], e["warnings"]) if part) or "无"
            detail.append([i, e["name"], e["status"], e["sample_rows"], "；".join(f"{label}: {count}" for label, count in zip(labels, e["counts"])), message, e["path"], e["selection"], e["footer_check"], e.get("sha256", ""), e.get("source_columns", "")])
        else:
            detail.append([i, "", "缺少文件", None, "", "本轮未提供 CSV，Sheet1 对应行留空。", ""])
    for col, width in zip("ABCDEFGHIJK", [12, 63, 14, 14, 34, 75, 90, 65, 70, 70, 80]):
        detail.column_dimensions[col].width = width
    for row in detail:
        for cell in row:
            cell.font = Font(name="宋体", size=11)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        detail.row_dimensions[row[0].row].height = 75 if row[0].row > 1 else 32
    detail.freeze_panes = "C2"
    detail.auto_filter.ref = "A1:K11"
    _save_with_caches(book, output, cache)
    return {"path": str(output), "input_dir": str(folder), "file_count": len(entries), "imported_files": imported,
            "skipped_files": len(entries)-imported, "region_count": len(labels), "labels": labels}


def write_folder_markdown(folder, output, schema, entries):
    """The same unrounded per-file means and fixed-FS formulas as Sheet1."""
    labels, columns, meaning = schema
    imported = [e for e in entries if e['status'] == '已导入']
    average = [mean(e['means'][j] for e in imported) for j in range(len(labels))]
    deviations = [[abs(e['means'][j]-average[j])/FS*100 for j in range(len(labels))] for e in imported]
    maxima = [max(row[j] for row in deviations) for j in range(len(labels))]
    averages = [mean(row[j] for row in deviations) for j in range(len(labels))]
    escape = lambda value: str(value).replace('|', '\\|').replace('\r', ' ').replace('\n', ' ')
    lines = [f'# {escape(Path(folder).name)} FS一致性结果', '',
             f'来源：{escape(folder)}', '', f'{meaning}；FS固定30 N。已导入 {len(imported)}/10 轮。', '',
             '每个CSV取全部有效采样的区域算术均值，各轮等权。读取CSV中已保存的区域值，不重新套用当前校正参数。', '',
             '| 轮次 | '+' | '.join(labels)+' |', '|---|'+'---:|'*len(labels)]
    for e in imported:
        lines.append(f"| {e['round']} | "+' | '.join(f'{v:.6f}' for v in e['means'])+' |')
    for label, values, unit in [('各轮均值（N）',average,''), ('FS%MAX',maxima,'%'), ('FS%mean',averages,'%')]:
        lines.append('| '+label+' | '+' | '.join(f'{v:.6f}{unit}' for v in values)+' |')
    lines += ['', f'八区/各区最差 FS%MAX：{max(maxima):.6f}%；最差 FS%mean：{max(averages):.6f}%。', '',
              '公式：单轮区域均值 mᵢⱼ=该轮该区有效值算术平均；区域中心 μⱼ=各轮 mᵢⱼ 的算术平均。', '',
              'FS%MAXⱼ=maxᵢ|mᵢⱼ−μⱼ|÷30×100%；FS%meanⱼ=meanᵢ|mᵢⱼ−μⱼ|÷30×100%。', '',
              '这是围绕本组均值的轮间一致性，不能代替相对测力计真值的准确度。H3区域值为估计分配，尚无各区独立真值。', '',
              '## 数据核对', '', '| 轮次 | 文件 | 有效行数 | 提示 | SHA256 |', '|---:|---|---:|---|---|']
    for e in entries:
        info='；'.join(str(e.get(k,'')) for k in ('error','warnings') if e.get(k)) or '无'
        lines.append(f"| {e['round']} | {escape(e['name'])} | {e['sample_rows']} | {escape(info)} | {e.get('sha256','')} |")
    if any(e.get('baseline_audit',{}).get('rejected') for e in entries):
        lines += ['','**存在基线质量异常：本表保留全部采样计算FS，数值满足阈值也不能代替本轮质量验收，应重新采集复核。**','']
    if len(imported) != 10:
        lines += ['','**不足10轮：仅作现有数据统计，不构成完整10轮复测。单轮会得到0偏差，不能据此判断一致性合格。**','']
    output=Path(output);output.parent.mkdir(parents=True,exist_ok=True)
    output.write_text('\n'.join(lines)+'\n',encoding='utf8')
    return dict(path=str(output),input_dir=str(folder),file_count=len(entries), imported_files=len(imported),
                skipped_files=len(entries)-len(imported),region_count=len(labels),labels=labels,
                fs_max_pct=maxima,fs_mean_pct=averages,output_format='md')


def generate_folder_reports(input_dir, output_dir, progress=None, output_format='xlsx'):
    if output_format not in ('xlsx','md'):
        raise ValueError('区域报告格式仅支持 xlsx 或 md。')
    root = Path(str(input_dir).strip().strip('"'))
    paths = csv_paths(root)
    folders = sorted({p.parent for p in paths}, key=lambda p: str(p))
    output_dir = Path(output_dir)
    if output_dir.suffix.lower() == ".xlsx":
        raise ValueError("按文件夹生成时请选择输出目录，不是单个 .xlsx 文件。")
    output_dir.mkdir(parents=True, exist_ok=True)
    reports, errors = [], []
    base = root.parent if root.is_file() else root
    for index, folder in enumerate(folders):
        if progress:
            progress(int(index / len(folders)*95)+1, f"统计 {folder.name}（一文件一轮）")
        relative = folder.relative_to(base)
        destination = output_dir / relative / f"{folder.name}_测试数据结果分析表.{output_format}"
        try:
            schema, entries = load_folder(folder)
            writer = write_folder_markdown if output_format == 'md' else write_folder_report
            reports.append(writer(folder, destination, schema, entries))
        except (ValueError, OSError) as exc:
            errors.append(f"{folder}：{exc}")
    if not reports:
        raise ValueError("没有生成结果表。\n" + "\n".join(errors))
    if errors:
        (output_dir / "未生成文件夹说明.txt").write_text("\n".join(errors), encoding="utf-8-sig")
    if progress:
        progress(100, f"已生成 {len(reports)} 张结果表")
    return {"report_mode": "folders", "output_format": output_format, "path": str(output_dir), "reports": reports, "errors": errors,
            "report_count": len(reports), "file_count": sum(r["imported_files"] for r in reports),
            "skipped_files": sum(r["skipped_files"] for r in reports), "failed_folders": len(errors)}

import csv
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from consistency_report import generate_consistency_report
from recording_report import acquisition_diagnostics, generate_recording_report, invalid_paired_capture, load_recordings, recording_groups


HEADER = ["record_type", "local_time_iso", "device_ms", "selected_resultant_or_diagnostic_n",
          "selected_value_mode", "stable_3s_sample", "measured_reference_resultant_n",
          "frame_state", "quality_flags", "position_mark", "note"]


def write_recording(folder, position="前", force=3, trial=1, *, values=None, mode="H3 Fz 候选模型", rows=None):
    mark = f"{position}{force}N_{trial:02}"
    path = Path(folder) / f"{mark}.csv"
    if rows is None:
        values = values if values is not None else [force - 0.1, force, force + 0.1]
        rows = [["PAIRED", f"2026-10-13T12:{force:02}:{trial:02}.{i:03}{position}", i * 20,
                 value, mode, 0, "", "PAIRED", "MODEL_VALIDATION_PENDING", mark, ""] for i, value in enumerate(values)]
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(HEADER)
        writer.writerows(rows)
    return path


class RecordingReportTests(unittest.TestCase):
    def test_continuous_trajectory_is_rejected_before_filename_condition(self):
        with tempfile.TemporaryDirectory() as directory:
            path = write_recording(directory, force=25, values=[25, 25])
            with path.open(encoding="utf-8-sig", newline="") as stream:
                rows = list(csv.reader(stream))
            rows[0].append("test_protocol")
            for row in rows[1:]:
                row.append("high_force_continuous")
            with path.open("w", encoding="utf-8-sig", newline="") as stream:
                csv.writer(stream).writerows(rows)
            trials, sources = load_recordings(path, allow_empty=True)
            self.assertEqual(trials, [])
            self.assertEqual(sources[0]["reason_code"], "continuous_protocol")
            self.assertIn("第 2 行", sources[0]["detail"])
            self.assertIn("多档加载/卸载", sources[0]["detail"])
            self.assertIn("按平台", sources[0]["action"])

    def test_acquisition_diagnostics_are_visible_and_not_measurements(self):
        with tempfile.TemporaryDirectory() as directory:
            path = write_recording(directory, values=[2, 4])
            with path.open(encoding="utf-8-sig", newline="") as stream:
                rows = list(csv.reader(stream))
            rows[0] += ["firmware_diagnostic_id", "firmware_error_count"]
            rows[1] += ["test-fw", 8]
            rows[2] += ["test-fw", 10]
            event = ["DIAG_INVALID", "", "", 999, "mode", 0, "", "INVALID", "ACQUISITION_INVALID", "前3N_01", "", "test-fw", 10]
            rows.insert(2, event)
            with path.open("w", encoding="utf-8-sig", newline="") as stream:
                csv.writer(stream).writerows(rows)
            trials, sources = load_recordings(path)
            self.assertEqual(trials[0]["used_frames"], 2)
            self.assertEqual(trials[0]["mean"], 3)
            self.assertEqual(trials[0]["invalid_cycles"], 1)
            self.assertIn("CSV 行：3", sources[0]["detail"])
            self.assertIn("8→10", sources[0]["detail"])
            self.assertIn("首末差：2", sources[0]["detail"])
            output = Path(directory)/"report.xlsx"
            generate_recording_report(path, output)
            book = load_workbook(output)
            self.assertIn("明确无效采集周期", book["重复采集评价"]["A22"].value)
            self.assertIn("test-fw", book["导入说明"]["D5"].value)
            book.close()

    def test_invalid_only_source_retains_diagnostic_reason(self):
        with tempfile.TemporaryDirectory() as directory:
            path = write_recording(directory, rows=[["DIAG_INVALID", "", "", "", "mode", 0, "", "INVALID", "ACQUISITION_INVALID", "前3N_01", ""]])
            trials, sources = load_recordings(path, allow_empty=True)
            self.assertEqual(trials, [])
            self.assertEqual(sources[0]["invalid_cycles"], 1)
            self.assertIn("无效采集周期 1 条", sources[0]["detail"])
            self.assertIn("没有 PAIRED", sources[0]["detail"])

    def test_unavailable_diagnostics_do_not_drop_legacy_firmware_values(self):
        with tempfile.TemporaryDirectory() as directory:
            path = write_recording(directory, values=[2, 3, 4])
            source = path.read_text(encoding="utf-8-sig")
            path.write_text(source.replace("MODEL_VALIDATION_PENDING", "MODEL_VALIDATION_PENDING|DIAG_UNAVAILABLE"), encoding="utf-8-sig")
            trials, _ = load_recordings(path)
            self.assertEqual(trials[0]["used_frames"], 3)
            self.assertEqual(trials[0]["mean"], 3)
            self.assertIn("固件诊断不可用", trials[0]["note"])

    def test_cumulative_failure_count_not_misreported_as_invalid_cycles(self):
        result = acquisition_diagnostics([
            {"record_type": "PAIRED", "firmware_error_count": "10"},
            {"record_type": "DIAG_ERROR_COUNT", "firmware_error_count": "13"},
            {"record_type": "PAIRED", "firmware_error_count": "13"},
        ])
        self.assertEqual(result["invalid_cycles"], 0)
        self.assertEqual(result["event_count"], 1)
        self.assertIn("首末差：3", result["text"])
        self.assertIn("非无效帧数", result["text"])

    def test_firmware_diagnostic_mode_does_not_require_host_tare(self):
        self.assertFalse(invalid_paired_capture({
            "record_type": "PAIRED", "selected_value_mode": "固件原始 Fz 分量",
            "host_tare_valid": "0", "acquisition_diag_available": "1",
            "firmware_valid_mask": "255", "firmware_baseline_epoch": "1",
        }))

    def test_exported_spatial_parameters_reproduce_eight_recorded_values(self):
        import numpy as np
        from PyQt5 import QtWidgets
        from main import Monitor, spatial_map_for_total
        app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])
        with tempfile.TemporaryDirectory() as directory:
            monitor = Monitor()
            try:
                monitor.raw = np.arange(24, dtype=float)*300 - 1800
                monitor.tare_raw = np.full(24, 7.0)
                monitor.tare_noise_floor = np.arange(8, dtype=float)*15 + 20
                monitor.device_fz = 15.2
                monitor.data_stale = False
                monitor.reconstructed_force_map = spatial_map_for_total(monitor.raw, monitor.tare_raw, monitor.device_fz, monitor.tare_noise_floor)[0]
                monitor._render(100)
                monitor.records = [monitor._record_base("PAIRED", 100, 1791878400)]
                path = Path(directory)/"replay.csv"
                with patch.object(QtWidgets.QFileDialog, "getSaveFileName", return_value=(str(path), "CSV Files (*.csv)")):
                    monitor._export()
                with path.open(encoding="utf-8-sig", newline="") as stream:
                    row = next(csv.DictReader(stream))
                self.assertNotIn(None, row)
                self.assertEqual(row["spatial_allocation_method"], "H3_MAG_NORM_NOISE_FLOOR_P135_V1")
                magnitudes = np.array([sum(float(row[f"S{i}_retared_dB{axis}"])**2 for axis in "xyz")**0.5 for i in range(1,9)])
                strengths = np.maximum(magnitudes-np.array([float(row[f"S{i}_spatial_noise_floor"]) for i in range(1,9)]),0)**float(row["spatial_allocation_exponent"])
                replay = strengths / strengths.sum()*float(row["selected_resultant_or_diagnostic_n"])
                np.testing.assert_allclose(replay, [float(row[f"S{i}_resultant_allocation_estimated_n"]) for i in range(1,9)], rtol=1e-13)
            finally:
                monitor.close()

    def test_current_monitor_export_imports_without_summary(self):
        from PyQt5 import QtWidgets
        from main import Monitor
        app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])
        with tempfile.TemporaryDirectory() as directory:
            monitor = Monitor()
            try:
                monitor.position.setText("前15N_01")
                monitor.device_fz = 15.2
                monitor.tare_raw = monitor.raw.copy()
                monitor.data_stale = False
                monitor.quality_flags = ("MODEL_VALIDATION_PENDING",)
                monitor.records = [monitor._record_base("PAIRED", i * 20, 1791878400 + i * 0.02) for i in range(20)]
                path = Path(directory) / "TacAI_H3_resultant_20261013_120000.csv"
                with patch.object(QtWidgets.QFileDialog, "getSaveFileName", return_value=(str(path), "CSV Files (*.csv)")):
                    monitor._export()
                self.assertEqual(len(list(Path(directory).glob("*.csv"))), 1)
                trials, _ = load_recordings(path)
                self.assertEqual(trials[0]["force"], 15)
                self.assertEqual(trials[0]["used_frames"], 20)
                self.assertEqual(trials[0]["median"], 15.2)
            finally:
                monitor.close()

    def test_sixty_files_are_sixty_trials_not_frames(self):
        with tempfile.TemporaryDirectory() as directory:
            for position in ("前", "后"):
                for force in (3, 15, 25):
                    for trial in range(1, 11):
                        write_recording(directory, position, force, trial, values=[force + trial / 10] * 20)
            trials, sources = load_recordings(directory)
            self.assertEqual(len(trials), 60)
            groups = recording_groups(trials)
            self.assertTrue(all(group["count"] == 10 for group in groups))
            self.assertAlmostEqual(groups[0]["mean"], 3.55)
            self.assertAlmostEqual(groups[0]["std"], (0.825 / 9) ** 0.5)
            self.assertTrue(all(trial["reference"] is None for trial in trials))
            output = Path(directory) / "report.xlsx"
            result = generate_consistency_report(directory, output, mode="recordings")
            self.assertEqual(result["complete_groups"], 6)
            book = load_workbook(output)
            self.assertEqual(book.sheetnames, ["重复采集评价", "每次采集", "导入说明"])
            self.assertEqual(book["每次采集"].max_row, 64)
            self.assertIsNone(book["每次采集"]["K5"].value)

    def test_partial_batch_and_single_trial_std_are_missing_not_zero(self):
        with tempfile.TemporaryDirectory() as directory:
            path = write_recording(directory)
            trials, _ = load_recordings(path)
            groups = recording_groups(trials)
            self.assertEqual(groups[0]["status"], "缺 9 次")
            self.assertIsNone(groups[0]["std"])
            self.assertIsNone(groups[1]["mean"])
            self.assertEqual(groups[1]["count"], 0)

    def test_duplicates_and_cumulative_exports_are_not_independent_trials(self):
        with tempfile.TemporaryDirectory() as directory:
            first = write_recording(directory)
            (Path(directory) / "copy.csv").write_bytes(first.read_bytes())
            with first.open(encoding="utf-8-sig", newline="") as stream:
                rows = list(csv.reader(stream))[1:]
            for row in rows:
                row[9] = "前3N_02"
            write_recording(directory, trial=2, rows=rows)
            trials, sources = load_recordings(directory)
            self.assertEqual(len(trials), 1)
            self.assertEqual(sum(row["status"] == "未导入" for row in sources), 2)
            self.assertTrue(any("重复采样帧" in row["detail"] for row in sources))

    def test_stable_flags_quality_and_marks_filter_frames(self):
        with tempfile.TemporaryDirectory() as directory:
            base = ["PAIRED", "", "", 3, "mode", 1, 3.1, "PAIRED", "", "前3N_01", ""]
            rows = [base.copy() for _ in range(6)]
            rows[0][5] = 0
            rows[0][3] = 99
            rows[1][8] = "RAW_FAULT"
            rows[1][3] = 88
            rows[2][0] = "MARK"
            rows[2][3] = 77
            rows[3][7] = "STALE"
            rows[4][3] = 2.9
            rows[5][3] = 3.1
            path = write_recording(directory, rows=rows)
            trial = load_recordings(path)[0][0]
            self.assertEqual(trial["used_frames"], 2)
            self.assertEqual(trial["median"], 3)
            self.assertAlmostEqual(trial["measured_error"], -0.1)
            self.assertEqual(trial["selection"], "稳定采样标记帧")

    def test_different_conditions_in_one_file_are_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            rows = [["PAIRED", "", "", 3, "mode", 0, "", "PAIRED", "", mark, ""]
                    for mark in ("前3N_01", "前15N_01")]
            path = write_recording(directory, rows=rows)
            with self.assertRaisesRegex(ValueError, "一个 CSV 必须只含一次采集"):
                load_recordings(path)

    def test_mode_mixing_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            write_recording(directory, trial=1, mode="Fz model")
            write_recording(directory, trial=2, mode="Firmware diagnostic")
            with self.assertRaisesRegex(ValueError, "不同力值模式"):
                load_recordings(directory)

    def test_candidate_notice_is_retained_but_correction_settings_cannot_mix(self):
        with tempfile.TemporaryDirectory() as directory:
            first = write_recording(directory)
            trials, _ = load_recordings(first)
            self.assertEqual(trials[0]["used_frames"], 3)
            self.assertIn("模型待独立验证", trials[0]["note"])
            second = write_recording(directory, trial=2)
            text = second.read_text(encoding="utf-8-sig")
            second.write_text(text.replace("MODEL_VALIDATION_PENDING", "MODEL_VALIDATION_PENDING|HIGH_FORCE_CORRECTION_OFF"), encoding="utf-8-sig")
            with self.assertRaisesRegex(ValueError, "高力修正设置"):
                load_recordings(directory)

    def test_empty_and_unknown_formats_are_reported_without_renaming(self):
        with tempfile.TemporaryDirectory() as directory:
            write_recording(directory)
            (Path(directory) / "empty.csv").write_text(",".join(HEADER), encoding="utf-8-sig")
            (Path(directory) / "unknown.csv").write_text("a,b\n1,2\n", encoding="utf-8-sig")
            trials, sources = load_recordings(directory)
            self.assertEqual(len(trials), 1)
            self.assertTrue(any("只有表头" in row["detail"] for row in sources))
            self.assertTrue(any("缺少列" in row["detail"] for row in sources))


if __name__ == "__main__":
    unittest.main()

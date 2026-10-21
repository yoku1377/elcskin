import csv
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from folder_report import H3_COLUMNS, NO7_COLUMNS, generate_folder_reports, load_folder


def write_csv(folder, index, values, legacy=False, bad=False):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"record_{index:02}.csv"
    columns = [f"P{i:02}_Fz({i//3},{i%3})" for i in range(12)] if legacy else H3_COLUMNS
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["record_type", "quality_flags", "stable_3s_sample", *columns])
        for value in values:
            writer.writerow(["PAIRED", "MODEL_LOW_CONFIDENCE|MODEL_SOFT_GATE_RESCUE", 0,
                             *["" if bad and i == 0 else value*(i+1) for i in range(len(columns))]])
    return path


def write_config_csv(folder, index, metadata, *, columns=H3_COLUMNS):
    path = Path(folder) / f"record_{index:02}.csv"
    fields = list(dict.fromkeys(key for row in metadata for key in row))
    header = ["record_type", *columns, *(key for key in fields if key != "record_type")]
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=header)
        writer.writeheader()
        for index, values in enumerate(metadata, 1):
            writer.writerow({"record_type": "PAIRED", **dict.fromkeys(columns, index), **values})
    return path


class FolderReportTests(unittest.TestCase):
    def test_continuous_trajectory_cannot_be_averaged_as_one_trial(self):
        with tempfile.TemporaryDirectory() as temp:
            write_csv(temp, 1, [2, 4])
            write_config_csv(temp, 2, [{"test_protocol": "high_force_continuous"}] * 2)
            with self.assertRaisesRegex(ValueError, "record_02.csv.*第 2 行.*多档加载/卸载.*按平台"):
                load_folder(temp)

    def test_changed_single_file_configuration_rejects_group_not_just_bad_round(self):
        changes = {
            "selected_value_mode": ("H3", "firmware"),
            "scalar_model_status": ("MODEL", "MODEL_CORRECTION_DISABLED"),
            "spatial_allocation_method": ("norm", "different_method"),
            "spatial_allocation_exponent": ("1.35", "1.4"),
            "S3_spatial_noise_floor": ("20", "21"),
            "total_correction_signature": ("off", "model-sha256"),
            "spatial_correction_signature": ("", "model-sha256"),
        }
        for field, (first, second) in changes.items():
            with self.subTest(field=field), tempfile.TemporaryDirectory() as temp:
                write_csv(temp, 1, [2, 4])
                write_config_csv(temp, 2, [{field: first}, {field: second}])
                with self.assertRaisesRegex(ValueError, rf"record_02.csv.*{field}.*第 2 行.*第 3 行"):
                    load_folder(temp)

    def test_cross_file_model_mode_method_exponent_conflicts_have_both_sources(self):
        for field in ("selected_value_mode", "scalar_model_status", "spatial_allocation_method", "spatial_allocation_exponent"):
            with self.subTest(field=field), tempfile.TemporaryDirectory() as temp:
                first, second = ("1.35", "1.4") if field.endswith("exponent") else ("original", "changed")
                write_config_csv(temp, 1, [{field: first}])
                write_config_csv(temp, 2, [{field: second}])
                with self.assertRaisesRegex(ValueError, rf"{field}.*record_01.csv 第 2 行.*record_02.csv 第 2 行"):
                    load_folder(temp)

    def test_per_trial_tare_and_noise_floor_differences_remain_legal(self):
        with tempfile.TemporaryDirectory() as temp:
            for index in (1, 2):
                write_config_csv(temp, index, [{
                    "selected_value_mode": "H3", "scalar_model_status": "MODEL",
                    "spatial_allocation_method": "norm", "spatial_allocation_exponent": "1.35",
                    "host_tare_epoch": str(index), "S3_spatial_noise_floor": str(20 * index),
                }] * 2)
            _, entries = load_folder(temp)
            self.assertEqual([entry["status"] for entry in entries], ["已导入"] * 2)
            self.assertEqual([entry["means"] for entry in entries], [[1.5] * 8] * 2)

    def test_cross_file_correction_parameters_cannot_mix_under_same_base_model(self):
        for field in ("correction_signature", "total_correction_signature",
                      "spatial_correction_signature", "local_comp_correction_signature"):
            with self.subTest(field=field), tempfile.TemporaryDirectory() as temp:
                common = {"scalar_model_status": "SAME_BASE_MODEL",
                          "spatial_allocation_method": "same_method"}
                write_config_csv(temp, 1, [{**common, field: "coefficients-A"}])
                write_config_csv(temp, 2, [{**common, field: "coefficients-B"}])
                with self.assertRaisesRegex(ValueError, rf"{field}.*record_01.csv 第 2 行.*record_02.csv 第 2 行"):
                    load_folder(temp)

    def test_same_profile_bypass_samples_remain_part_of_the_trial(self):
        with tempfile.TemporaryDirectory() as temp:
            common = {"local_comp_correction_signature": "same-coefficients"}
            write_config_csv(temp, 1, [{**common, "local_comp_status": "APPLIED"},
                                       {**common, "local_comp_status": "BYPASS"}])
            write_config_csv(temp, 2, [{**common, "local_comp_status": "BYPASS"}])
            _, entries = load_folder(temp)
            self.assertEqual(entries[0]["sample_rows"], 2)
            self.assertEqual(entries[0]["means"], [1.5] * 8)
            self.assertEqual(entries[1]["sample_rows"], 1)

    def test_configuration_audit_ignores_events_and_invalid_paired(self):
        with tempfile.TemporaryDirectory() as temp:
            common = {"acquisition_diag_available": "1", "firmware_baseline_epoch": "1", "host_tare_valid": "1"}
            write_config_csv(temp, 1, [
                {**common, "selected_value_mode": "H3", "firmware_valid_mask": "255"},
                {**common, "selected_value_mode": "changed", "firmware_valid_mask": "254", "test_protocol": "high_force_continuous"},
                {**common, "record_type": "DIAG_RESET", "selected_value_mode": "event", "firmware_valid_mask": "255", "test_protocol": "high_force_continuous"},
                {**common, "selected_value_mode": "H3", "firmware_valid_mask": "255"},
            ])
            _, entries = load_folder(temp)
            self.assertEqual(entries[0]["sample_rows"], 2)
            self.assertEqual(entries[0]["means"], [2.5] * 8)

    def test_named_config_numeric_spelling_and_frame_values_are_not_changes(self):
        with tempfile.TemporaryDirectory() as temp:
            write_config_csv(temp, 1, [
                {"spatial_allocation_exponent": "1.35", "S1_spatial_noise_floor": "20", "correction_n": "0.1"},
                {"spatial_allocation_exponent": "1.3500", "S1_spatial_noise_floor": "2e1", "correction_n": "0.9"},
            ])
            write_config_csv(temp, 2, [{"spatial_allocation_exponent": "1.350000"}])
            _, entries = load_folder(temp)
            self.assertEqual(len(entries), 2)
            self.assertEqual(entries[0]["means"], [1.5] * 8)

    def test_nonfinite_config_cannot_silently_remove_a_trial(self):
        with tempfile.TemporaryDirectory() as temp:
            write_csv(temp, 1, [1])
            write_config_csv(temp, 2, [{"spatial_allocation_exponent": "NaN"}])
            with self.assertRaisesRegex(ValueError, "record_02.csv 第 2 行.*spatial_allocation_exponent.*有限数值"):
                load_folder(temp)

    def test_legacy_no7_raw_force_does_not_acquire_paired_metadata_requirements(self):
        with tempfile.TemporaryDirectory() as temp:
            write_config_csv(temp, 1, [
                {"record_type": "RAW", "selected_value_mode": "raw"},
                {"record_type": "FORCE", "selected_value_mode": "force"},
            ], columns=NO7_COLUMNS)
            _, entries = load_folder(temp)
            self.assertEqual(entries[0]["means"], [1.5] * 12)

    def test_different_host_tare_epochs_cannot_mix(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)/"record_01.csv"
            with path.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(["record_type", *H3_COLUMNS, "host_tare_epoch"])
                writer.writerow(["PAIRED", *([2]*8), 1])
                writer.writerow(["TARE_CHANGED", *([""]*8), 2])
                writer.writerow(["PAIRED", *([3]*8), 2])
            with self.assertRaisesRegex(ValueError, "host_tare_epoch.*置零基线变更"):
                load_folder(temp)

    def test_acquisition_events_never_become_values_and_remain_visible(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            path = root / "record_01.csv"
            with path.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(["record_type", "quality_flags", *H3_COLUMNS,
                                 "firmware_diagnostic_id", "firmware_error_count"])
                writer.writerow(["PAIRED", "OK", *([2]*8), "test-fw", 7])
                # Even a malformed producer that retains a value in an event
                # must never let it become a zero or a measurement.
                writer.writerow(["DIAG_INVALID", "ACQUISITION_INVALID", *([999]*8), "test-fw", 9])
                writer.writerow(["DIAG_ERROR_COUNT", "ACQUISITION_ERROR_COUNT_INCREASE", *([999]*8), "test-fw", 9])
                writer.writerow(["PAIRED", "OK", *([4]*8), "test-fw", 9])
            _, entries = load_folder(root)
            entry = entries[0]
            self.assertEqual(entry["means"], [3]*8)
            self.assertEqual(entry["sample_rows"], 2)
            self.assertEqual(entry["invalid_cycles"], 1)
            self.assertEqual(entry["acquisition_events"], 2)
            self.assertIn("7→9", entry["warnings"])
            self.assertIn("首末差：2", entry["warnings"])
            self.assertIn("CSV 行：3, 4", entry["warnings"])
            result = generate_folder_reports(root, root/"out")
            book = load_workbook(result["reports"][0]["path"], data_only=True)
            self.assertEqual(book["Sheet1"]["B3"].value, 3)
            self.assertEqual(book["逐帧区域值"].max_row, 3)
            self.assertIn("test-fw", book["输入明细"]["F2"].value)
            self.assertIn("仅覆盖有效记录", book["Sheet1"]["B3"].comment.text)
            self.assertIn("第 1 轮含采集异常", book["Sheet1"]["A58"].value)
            book.close()

    def test_invalid_paired_metadata_is_not_used_and_invalid_only_has_diagnostics(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)/"record_01.csv"
            header = ["record_type", *H3_COLUMNS, "acquisition_diag_available", "firmware_valid_mask", "firmware_baseline_epoch", "host_tare_valid", "selected_value_mode"]
            with path.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(header)
                writer.writerow(["PAIRED", *([2]*8), 1, "0xFF", 1, 1, "H3 V3-3 Fz 候选模型（0.2–30 N）"])
                writer.writerow(["PAIRED", *([99]*8), 1, "0xFE", 1, 1, "H3 V3-3 Fz 候选模型（0.2–30 N）"])
                writer.writerow(["PAIRED", *([100]*8), 1, "0xFF", 1, 0, "H3 V3-3 Fz 候选模型（0.2–30 N）"])
            _, entries = load_folder(temp)
            self.assertEqual(entries[0]["means"], [2]*8)
            self.assertEqual(entries[0]["invalid_cycles"], 0)
            self.assertEqual(entries[0]["acquisition_events"], 2)
            self.assertIn("诊断信息无效的配对记录 2 条", entries[0]["warnings"])
            with path.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(header)
                writer.writerow(["DIAG_INVALID", *([""]*8), 1, "0xFE", 1, 1, "H3 V3-3 Fz 候选模型（0.2–30 N）"])
            with self.assertRaisesRegex(ValueError, "无效采集周期 1 条"):
                load_folder(temp)

    def test_no7_raw_force_unnamed_columns_and_saved_mean_are_handled(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            folder = root / "前15N"
            folder.mkdir()
            with (folder / "NO7_01.csv").open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(["record_type", "device_ms", *NO7_COLUMNS, "", ""])
                writer.writerow(["RAW", 100, *range(1,13), "", 99])
                writer.writerow(["FORCE", 100, *range(3,15), "", 99])
                writer.writerow(["MARK", 100, *([99]*12), "", 99])
                writer.writerow([""]*16)
                writer.writerow(["", "", *range(2,14), "", 99])
            schema, entries = load_folder(folder)
            self.assertEqual(schema[1], NO7_COLUMNS)
            self.assertEqual(entries[0]["means"], list(range(2,14)))
            self.assertEqual(entries[0]["sample_rows"], 2)
            self.assertIn("一致", entries[0]["footer_check"])
            self.assertEqual(entries[0]["footer_max_error"], 0)
            result = generate_folder_reports(folder, root/"out")
            book = load_workbook(result["reports"][0]["path"], data_only=True)
            self.assertEqual(book["Sheet1"]["M3"].value, 13)
            self.assertIn("RAW 1 行", book["输入明细"]["H2"].value)
            self.assertIn("一致", book["输入明细"]["I2"].value)
            book.close()

    def test_no7_wrong_saved_mean_is_flagged_never_used_as_input(self):
        with tempfile.TemporaryDirectory() as temp:
            with (Path(temp)/"NO7_01.csv").open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(["record_type", *NO7_COLUMNS])
                writer.writerow(["RAW", *([1]*12)])
                writer.writerow(["FORCE", *([3]*12)])
                writer.writerow(["", *([50]*12)])
            _, entries = load_folder(temp)
            self.assertEqual(entries[0]["means"], [2]*12)
            self.assertIn("有差异", entries[0]["footer_check"])
            self.assertEqual(entries[0]["footer_max_error"], 48)

    def test_duplicate_named_columns_still_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            with (Path(temp)/"NO7_01.csv").open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(["record_type", *NO7_COLUMNS, NO7_COLUMNS[0]])
                writer.writerow(["RAW", *([1]*13)])
            with self.assertRaisesRegex(ValueError, "重复的非空列名"):
                load_folder(temp)

    def test_one_folder_one_workbook_and_all_numeric_frames_arithmetic_mean(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            for folder in ("前3N", "后15N"):
                for i in range(1, 11):
                    write_csv(root/"input"/folder, i, [i, i, i+9])
            result = generate_folder_reports(root/"input", root/"out")
            self.assertEqual(result["report_count"], 2)
            self.assertEqual(result["file_count"], 20)
            self.assertEqual(result["skipped_files"], 0)
            path = result["reports"][0]["path"]
            values = load_workbook(path, data_only=True)
            formulas = load_workbook(path, data_only=False)
            sheet = values["Sheet1"]
            self.assertEqual(values.sheetnames, ["Sheet1", "输入明细", "逐帧区域值"])
            self.assertEqual(sheet["B3"].value, 4)  # median would be 1
            self.assertEqual(sheet["I3"].value, 32)
            self.assertEqual(sheet["B15"].value, 8.5)
            self.assertEqual(sheet["B18"].value, 4.5)
            self.assertAlmostEqual(sheet["B30"].value, .15)
            self.assertAlmostEqual(sheet["B42"].value, .15)
            self.assertAlmostEqual(sheet["B44"].value, .15*36)
            self.assertAlmostEqual(sheet["B46"].value, 2.5/30)
            self.assertAlmostEqual(sheet["B48"].value, 3)
            self.assertIn("SUM(B42:I42)", formulas["Sheet1"]["B44"].value)
            self.assertEqual(formulas["Sheet1"]["B3"].value, "=AVERAGE('逐帧区域值'!F2:F4)")
            self.assertEqual(values["逐帧区域值"]["F4"].value, 10)
            self.assertEqual(values["逐帧区域值"]["B4"].value, 4)
            self.assertEqual(len(values["输入明细"]["J2"].value), 64)
            self.assertIn("保留数值", values["输入明细"]["F2"].value)
            self.assertTrue(all("合格" not in str(cell.value) for row in sheet for cell in row if cell.row < 50))
            values.close()
            formulas.close()

    def test_missing_and_bad_round_keep_blank_not_zero(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            write_csv(root/"input", 1, [0, 0])
            write_csv(root/"input", 3, [2, 2], bad=True)
            write_csv(root/"input", 4, [4, 4])
            result = generate_folder_reports(root/"input", root/"out")
            self.assertEqual(result["file_count"], 2)
            self.assertEqual(result["skipped_files"], 1)
            book = load_workbook(result["reports"][0]["path"], data_only=True)
            sheet = book["Sheet1"]
            self.assertEqual(sheet["B3"].value, 0)
            self.assertIsNone(sheet["B4"].value)
            self.assertIsNone(sheet["B5"].value)
            self.assertEqual(sheet["B6"].value, 4)
            self.assertEqual(sheet["B15"].value, 2)
            self.assertIsNone(sheet["B19"].value)
            self.assertEqual(book["输入明细"]["C3"].value, "缺少文件")
            self.assertEqual(book["输入明细"]["C4"].value, "未导入")
            book.close()

    def test_legacy_twelve_zone_input_preserves_template_twelve_columns(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            write_csv(root/"input", 1, [1, 2, 3], legacy=True)
            result = generate_folder_reports(root/"input", root/"out")
            book = load_workbook(result["reports"][0]["path"], data_only=True)
            self.assertEqual(book["Sheet1"]["M2"].value, "P11_Fz(mean)")
            self.assertEqual(book["Sheet1"]["M3"].value, 24)
            book.close()

    def test_eleven_files_never_silently_drop_last_file(self):
        with tempfile.TemporaryDirectory() as temp:
            for i in range(1, 12):
                write_csv(temp, i, [1, 2])
            with self.assertRaisesRegex(ValueError, "11 个 CSV"):
                load_folder(temp)

    def test_mark_rows_not_measurements_and_nan_is_reported(self):
        with tempfile.TemporaryDirectory() as temp:
            path = write_csv(temp, 1, [1, float('nan'), 3])
            with path.open("a", encoding="utf-8", newline="") as stream:
                csv.writer(stream).writerow(["MARK", "", 0, *([999]*8)])
            _, entries = load_folder(temp)
            self.assertEqual(entries[0]["means"][0], 2)
            self.assertEqual(entries[0]["counts"][0], 2)
            self.assertIn("非有限", entries[0]["warnings"])

    def test_trace_preserves_csv_row_numbers_and_named_column_positions(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)/"record_01.csv"
            with path.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(["record_type", "", *H3_COLUMNS])
                writer.writerow(["MARK", "", *([99]*8)])
                writer.writerow([""]*10)
                writer.writerow(["PAIRED", "", *([2]*8)])
                writer.writerow(["PAIRED", "", "NaN", *([4]*7)])
            result = generate_folder_reports(temp, Path(temp)/"out")
            book = load_workbook(result["reports"][0]["path"], data_only=True)
            self.assertEqual(book["逐帧区域值"]["B2"].value, 4)
            self.assertIsNone(book["逐帧区域值"]["F3"].value)
            self.assertEqual(book["Sheet1"]["B3"].value, 2)
            self.assertEqual(book["Sheet1"]["C3"].value, 3)
            self.assertIn("S1=C列", book["输入明细"]["K2"].value)
            book.close()

    def test_same_leaf_folder_names_never_overwrite(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            write_csv(root/"input"/"设备A"/"前3N", 1, [1, 1])
            write_csv(root/"input"/"设备B"/"前3N", 1, [2, 2])
            result = generate_folder_reports(root/"input", root/"out")
            self.assertEqual(len({r["path"] for r in result["reports"]}), 2)


if __name__ == "__main__":
    unittest.main()

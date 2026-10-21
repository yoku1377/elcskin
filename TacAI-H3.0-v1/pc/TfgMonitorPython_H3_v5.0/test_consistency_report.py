import csv
import os
from pathlib import Path
import tempfile
import unittest

os.environ.setdefault("QT_QPA_PLATFORM", "windows")

from openpyxl import load_workbook
from PyQt5 import QtGui, QtWidgets

from consistency_report import generate_consistency_report, read_recoverable_csv, load_sessions, build_metrics
from consistency_report_dialog import ConsistencyReportDialog


SUMMARY_HEADER = [
    "completed_time_iso", "device_sn", "step_index", "repeat", "position", "direction",
    "target_n", "reference_median_n", "temperature_c", "upper_model_median_n",
    "firmware_fz_median_n", "raw_activity_median", "quality",
]


def write_session(folder: Path, device: str, stamp: str, offset: float = 0.0):
    summary = folder / f"H3_calibration_{device}_{stamp}_summary.csv"
    frames = folder / f"H3_calibration_{device}_{stamp}_frames.csv"
    rows = []
    step = 0
    for repeat in (1, 2):
        for position in ("前", "中", "后"):
            for direction, references in (("加载", (0.0, 0.2, 1.0, 5.0)), ("卸载", (5.0, 1.0, 0.2, 0.0))):
                for reference in references:
                    step += 1
                    output = max(0.0, reference + offset + (0.02 if repeat == 2 else 0.0))
                    rows.append([
                        "2026-10-08T12:00:00", device, step, repeat, position, direction,
                        reference, reference, 25.0, output, output, 1000.0, "OK",
                    ])
    with summary.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(SUMMARY_HEADER)
        writer.writerows(rows)
    with frames.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(["local_time_iso", "device_sn", "silicone_batch", "magnet_assembly", "operator"])
        writer.writerow(["2026-10-08T12:00:00", device, "SIL-01", "MAG-01", "tester"])


class ConsistencyReportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])

    def test_single_device_report_and_empty_session_skip(self):
        with tempfile.TemporaryDirectory() as folder_name:
            folder = Path(folder_name)
            write_session(folder, "H3-001", "20261008_120000")
            (folder / "H3_calibration_empty_20261008_120100_summary.csv").write_text(
                ",".join(SUMMARY_HEADER) + "\n", encoding="utf-8-sig"
            )
            output = folder / "report.xlsx"
            result = generate_consistency_report(folder, output)
            self.assertEqual(result["device_count"], 1)
            self.assertEqual(result["skipped_empty"], 1)
            workbook = load_workbook(output, data_only=False)
            self.assertEqual(
                workbook.sheetnames,
                ["综合评价", "判定参数", "设备指标", "位置与力点", "设备间一致性", "原始步汇总", "来源说明"],
            )
            self.assertEqual(workbook["设备指标"]["A5"].value, "H3-001")
            self.assertTrue(str(workbook["设备指标"]["U5"].value).startswith("=IF("))
            self.assertIn("尚缺第二台设备", workbook["综合评价"]["B8"].value)

    def test_two_devices_create_consistency_rows_and_rank_formulas(self):
        with tempfile.TemporaryDirectory() as folder_name:
            folder = Path(folder_name)
            write_session(folder, "H3-001", "20261008_120000", 0.0)
            write_session(folder, "H3-002", "20261008_120100", 0.15)
            output = folder / "report.xlsx"
            result = generate_consistency_report(folder, output)
            self.assertEqual(result["device_count"], 2)
            workbook = load_workbook(output, data_only=False)
            self.assertEqual(workbook["设备指标"]["A6"].value, "H3-002")
            self.assertIn("RANK(", str(workbook["设备指标"]["W5"].value))
            self.assertIsNotNone(workbook["设备间一致性"]["A5"].value)

    def test_legacy_header_shift_is_recovered(self):
        with tempfile.TemporaryDirectory() as folder_name:
            path = Path(folder_name) / "legacy_summary.csv"
            header = ["upper_model_median_n", "firmware_fz_median_n", "quality"]
            with path.open("w", newline="", encoding="utf-8-sig") as stream:
                writer = csv.writer(stream)
                writer.writerow(header)
                writer.writerow(["1.0", "0.9", "1", "1.1", "OK"])
            rows, recovered = read_recoverable_csv(path, "summary")
            self.assertTrue(recovered)
            self.assertEqual(rows[0]["upper_model_load_score_median"], "0.9")
            self.assertEqual(rows[0]["upper_model_soft_gate_rescue_frames"], "1")
            self.assertEqual(rows[0]["firmware_fz_median_n"], "1.1")

    def test_dialog_exposes_clear_generation_action(self):
        dialog = ConsistencyReportDialog()
        buttons = {button.text() for button in dialog.findChildren(QtWidgets.QPushButton)}
        self.assertIn("生成 Markdown 报告", buttons)
        self.assertIn("选择目录", buttons)
        self.assertIn("选择 CSV", buttons)
        self.assertEqual(dialog.mode_box.currentData(), "folders_md")
        self.assertFalse(dialog._default_output("D:/data").endswith(".xlsx"))
        dialog.mode_box.setCurrentIndex(1)
        self.assertTrue(dialog._default_output("D:/data").endswith(".xlsx"))
        dialog.mode_box.setCurrentIndex(0)
        screenshot = os.environ.get("H3_DIALOG_SCREENSHOT")
        if screenshot:
            dialog.resize(dialog.sizeHint())
            image = QtGui.QImage(dialog.size(), QtGui.QImage.Format_ARGB32)
            image.fill(QtGui.QColor("#061018"))
            painter = QtGui.QPainter(image)
            dialog.render(painter)
            painter.end()
            self.assertTrue(image.save(screenshot))
        dialog.close()

    def test_direct_summary_file_and_missing_frames(self):
        with tempfile.TemporaryDirectory() as folder_name:
            folder = Path(folder_name)
            write_session(folder, "H3-001", "20261008_120000")
            original = next(folder.glob("*_summary.csv"))
            original.rename(folder / "summary.csv")
            rows, sources, metadata = load_sessions(folder / "summary.csv")
            self.assertEqual(len(rows), 48)
            self.assertEqual(sources[0]["frame_rows"], 0)

    def test_header_only_file_has_actionable_message(self):
        with tempfile.TemporaryDirectory() as folder_name:
            path = Path(folder_name) / "summary.csv"
            path.write_text(",".join(SUMMARY_HEADER), encoding="utf-8-sig")
            with self.assertRaisesRegex(ValueError, "0 条采样记录"):
                load_sessions(path)

    def test_recording_csv_renamed_summary_does_not_become_calibration(self):
        with tempfile.TemporaryDirectory() as folder_name:
            path = Path(folder_name) / "summary.csv"
            path.write_text("record_type,position_mark\nPAIRED,前3N_01\n", encoding="utf-8-sig")
            with self.assertRaisesRegex(ValueError, "不是标定汇总格式"):
                load_sessions(path)

    def test_partial_high_force_data_has_missing_metrics_and_guarded_status(self):
        with tempfile.TemporaryDirectory() as folder_name:
            folder = Path(folder_name)
            write_session(folder, "H3-001", "20261008_120000")
            path = next(folder.glob("*_summary.csv"))
            with path.open(encoding="utf-8-sig", newline="") as stream:
                rows = list(csv.reader(stream))
            with path.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(rows[0])
                writer.writerows(row for row in rows[1:] if row[4] != "中" and row[5] == "加载" and row[6] == "5.0")
            loaded, sources, metadata = load_sessions(path)
            metric = build_metrics(loaded, metadata)["device_metrics"][0]
            for key in ("low_force_mae_n", "hysteresis_p95_n", "residual_median_n", "position_spread_p95_n"):
                self.assertIsNone(metric[key])
            output = folder / "partial.xlsx"
            generate_consistency_report(path, output)
            book = load_workbook(output)
            self.assertIn("COUNT(", book["设备指标"]["U5"].value)
            self.assertIn("数据不足", book["设备指标"]["V5"].value)
            self.assertIn("ISNUMBER", book["设备指标"]["W5"].value)


if __name__ == "__main__":
    unittest.main()

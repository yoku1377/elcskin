"""One recording CSV is one trial; frames never count as independent trials."""

from __future__ import annotations

import csv
import hashlib
import io
import math
import re
import unicodedata
from collections import Counter
from pathlib import Path
from statistics import mean, median, stdev

from openpyxl import Workbook
from openpyxl.styles import Alignment

CONDITIONS = [(position, force) for position in ("前", "后") for force in (3, 15, 25)]
EXPECTED_TRIALS = 10
VALUE_COLUMN = "selected_resultant_or_diagnostic_n"
REQUIRED_COLUMNS = {"record_type", VALUE_COLUMN, "selected_value_mode"}
INFORMATION_FLAGS = {"OK", "MODEL_VALIDATION_PENDING", "HIGH_FORCE_CORRECTION_OFF", "DIAG_UNAVAILABLE", "BASELINE_QUALITY_PENDING"}
FLAG_LABELS = {
    "RELEASE_BASELINE_COMPENSATED": "工作零点已补偿；本记录不能证明传感器自然回零",
    "RELEASE_BASELINE_UPDATED": "松手后工作零点更新；须按同一工作零点分别评价",
    "LOCAL_CORRECTION_NOT_APPLIED": "已选局部补偿但高力帧未应用，需核对旁路原因（原始值保留）",
    "BASELINE_QUALITY_PENDING": "静置基线尚未就绪（原始数值保留）",
    "BASELINE_QUALITY_REJECTED": "静置基线检查未通过，专项补偿已旁路（原始数值保留）",
    "MODEL_VALIDATION_PENDING": "模型待独立验证",
    "HIGH_FORCE_CORRECTION_OFF": "高力修正已关闭",
    "LOW_SIGNAL_THRESHOLD": "信号接近空载阈值",
    "MODEL_LOW_CONFIDENCE": "模型对当前读数把握较低",
    "MODEL_SOFT_GATE_RESCUE": "分类判断异常，已保留连续读数",
    "MODEL_DIRECTION_INVARIANT_FALLBACK": "使用斜压或未知方向估计值",
    "MODEL_CLASSIFIER_IGNORED_CONTINUITY": "分类判断与读数不一致",
    "MODEL_KNOWN_RELEASE_ZERO": "检测到卸载，读数已归零",
    "MODEL_REGRESSION_NONPOSITIVE": "模型原始估计值不大于零",
    "MAGNETIC_COHERENCE": "可能存在磁干扰",
    "UPPER_MODEL_SATURATED": "上位机读数达到量程上限",
    "FIRMWARE_SATURATED": "固件读数达到量程上限",
    "UNCALIBRATED_POSITION": "当前按压位置未经标定",
    "DIAG_UNAVAILABLE": "固件诊断不可用，无法核实采集错误计数",
}
MEASUREMENT_WARNINGS = FLAG_LABELS.keys() - INFORMATION_FLAGS
ACQUISITION_EVENTS = {
    "DIAG_INVALID": "无效采集周期",
    "DIAG_ERROR_COUNT": "失败通道累计计数增长",
    "DIAG_RESET": "采集基线或设备状态重置",
    "STREAM_ERROR": "数据流错误",
    "DATA_TIMEOUT": "数据超时",
    "STREAM_DISCONNECTED": "数据流断开",
    "RAW_NO_TARE": "尚未置零的原始采样",
    "TARE_CHANGED": "上位机置零基线变更",
}
PROBLEMS = {
    "empty": ("文件没有采样数据", "换选实际记录过数据的 CSV。"),
    "format": ("CSV 格式未识别", "提供这个原始 CSV 核对格式，不要只修改文件名。"),
    "mark": ("无法确定按压位置和目标力", "位置标记或文件名填写 前15N_01 这样的工况。"),
    "mixed_condition": ("一个文件包含多个工况或采集编号", "按一次采集一个文件导出，导出后清空本轮。"),
    "continuous_protocol": ("高力连续开发序列不能作为单轮评价", "保留完整原始记录；按平台明确切分并保留来源行号、时间和区间规则后评价。"),
    "conflict": ("文件名与记录的工况不一致", "核对前/后位置和目标力，修正错误标记后重试。"),
    "condition": ("不属于这次六组测试", "请选择前/后 3、15、25 N 的记录。"),
    "duplicate": ("同一份数据重复导出", "保留其中一份即可，复制文件不能增加采集次数。"),
    "overlap": ("与另一文件包含相同采样", "核对是否未清空本轮就再次导出；同一段采样不重复计数。"),
    "mode": ("文件内的测量设置不一致", "一次记录使用相同力值模式和高力修正设置。"),
    "device": ("一个文件包含多台设备", "按设备分别导出。"),
    "no_frames": ("没有可识别的采样记录", "提供原始文件核对导出格式；仅有位置标记不算采样。"),
    "too_few": ("可用采样点不足", "查看被排除的具体原因；确认设备有实时读数并完成空载置零后重新采集。"),
    "read": ("文件无法读取", "确认文件存在且可读取；把原始文件提供给开发核对。"),
}


class ImportProblem(ValueError):
    def __init__(self, code, detail):
        super().__init__(detail)
        self.code = code
MARK_PATTERN = re.compile(r"(前|后|front|rear|back)\s*(?:区)?\s*[_ -]?\s*(\d+(?:\.\d+)?)\s*N(?:[_ -]+(\d+))?(?![\d.])", re.I)


def _finite(value):
    try:
        number = float(value)
        return number if math.isfinite(number) else None
    except (ValueError, TypeError):
        return None


def _flags(row):
    return {flag.strip().upper() for flag in row.get("quality_flags", "").split("|") if flag.strip()}


def audit_effective_baseline(rows):
    """A fixed working zero is legal; changing it within a trial is not.

    Legacy CSVs do not have a working-zero epoch. Invalid acquisition events
    cannot establish a conflicting coordinate system or prove a valid one.
    Epochs are local to a trial and must never be compared across files.
    """
    samples = [row for row in rows
               if row.get("record_type", "PAIRED").strip().upper() == "PAIRED"
               and not invalid_paired_capture(row)]
    field = "effective_baseline_epoch"
    current_schema = any(
        (match := re.match(r"^5\.3\.(\d+)(?:\D|$)", str(row.get("monitor_version", ""))))
        and int(match[1]) >= 23 for row in samples
    )
    if not samples or not (current_schema or any(field in row for row in samples)):
        return ""
    epochs = set()
    for row in samples:
        value = _finite(row.get(field))
        if value is None or value < 0 or value != int(value):
            raise ValueError(f"{field} 工作零点记录缺失或无效；不能核实同一基线，不能混合评价。")
        epochs.add(int(value))
    if len(epochs) != 1:
        raise ValueError(f"同一文件采样跨多个工作零点（{field}）；请按工作零点分别导出，不能混合评价。")
    offset_fields = [f"S{i}_release_offset_dB{axis}" for i in range(1, 9) for axis in "xyz"]
    have_offsets = any(any(field in row for field in offset_fields) for row in samples)
    compensated = False
    if have_offsets or current_schema:
        reference = None
        for row in samples:
            offset = tuple(_finite(row.get(field)) for field in offset_fields)
            if any(value is None for value in offset):
                raise ValueError("工作零点 RAW 补偿记录缺失或无效；不能重放该基线，不能混合评价。")
            if reference is not None and offset != reference:
                raise ValueError("同一 effective_baseline_epoch 内 RAW 工作零点补偿发生变化；不能混合评价。")
            reference = offset
            compensated = compensated or any(value != 0. for value in offset)
    if compensated or any("RELEASE_BASELINE_COMPENSATED" in _flags(row) for row in samples):
        return "工作零点已在本轮前补偿并于本记录内保持；本记录不能证明传感器自然回零"
    return ""


def _diagnostic_integer(value):
    text = str(value or "").strip()
    try:
        return int(text, 16) if text.lower().startswith("0x") else int(text)
    except ValueError:
        return None


def invalid_paired_capture(row):
    """Reject explicitly invalid acquisition metadata, never infer it from force."""
    if row.get("record_type", "PAIRED").strip().upper() != "PAIRED":
        return False
    # The firmware diagnostic force mode does not require a host RAW tare.
    requires_tare = "候选模型" in row.get("selected_value_mode", "")
    if requires_tare and str(row.get("host_tare_valid", "")).strip().lower() in {"0", "false"}:
        return True
    if str(row.get("acquisition_diag_available", "")).strip().lower() in {"1", "true"}:
        return (_diagnostic_integer(row.get("firmware_valid_mask")) != 255
                or (_diagnostic_integer(row.get("firmware_baseline_epoch")) or 0) <= 0)
    return False


def continuous_protocol_reason(rows):
    """A development trajectory cannot be averaged as one nominal-force trial."""
    for row in rows:
        if (row.get("record_type", "PAIRED").strip().upper() == "PAIRED"
                and not invalid_paired_capture(row)
                and row.get("test_protocol", "").strip().lower() == "high_force_continuous"):
            return (
                f"第 {row.get('_source_line', '?')} 行 test_protocol=high_force_continuous："
                "该文件是含多档加载/卸载的高力补偿开发连续采集资料，不能作为一个 25 N 单轮 CSV 求均值；"
                "请按平台和事先确定的区间规则明确切分，保留源文件、原始行号与时间范围后评价。")
    return None


def acquisition_diagnostics(rows):
    """Summarize loss separately from numerical samples and cumulative counters."""
    events = Counter()
    event_lines = []
    reasons = Counter()
    invalid_paired = 0
    for line, row in enumerate(rows, 2):
        kind = row.get("record_type", "").strip().upper()
        if invalid_paired_capture(row):
            invalid_paired += 1
            kind = "PAIRED_INVALID"
        if kind in ACQUISITION_EVENTS or kind == "PAIRED_INVALID":
            events[kind] += 1
            event_lines.append(str(row.get("_source_line", line)))
            reasons.update(_flags(row) - {"OK"})
    # Only an explicit firmware invalid-cycle event proves this count. A
    # malformed PAIRED row or missing host tare is not another firmware cycle.
    invalid_cycles = events["DIAG_INVALID"]
    messages = []
    if events:
        labels = {**ACQUISITION_EVENTS, "PAIRED_INVALID": "诊断信息无效的配对记录"}
        messages.append("采集异常：" + "、".join(f"{labels[kind]} {count} 条" for kind, count in sorted(events.items())))
        messages.append("上述事件未计入力值统计，未补零或沿用旧值；结果仅覆盖有效记录")
        messages.append("事件所在 CSV 行：" + ", ".join(event_lines[:20]) + (f" 等共 {len(event_lines)} 行" if len(event_lines) > 20 else ""))
        if reasons:
            messages.append("诊断标记：" + "、".join(f"{flag}({count})" for flag, count in sorted(reasons.items())))
    versions = sorted({row.get("firmware_diagnostic_id", "").strip() for row in rows} - {""})
    if versions:
        messages.append("固件诊断标识：" + "、".join(versions))
    counters = [_diagnostic_integer(row.get("firmware_error_count")) for row in rows]
    counters = [value for value in counters if value is not None]
    if counters:
        messages.append(f"固件累计失败通道数：{counters[0]}→{counters[-1]}（启动以来累计，可含录制前错误，不等于无效周期数）")
        if any(current < previous for previous, current in zip(counters, counters[1:])):
            messages.append("累计计数在录制期间回退，请核对设备重启或诊断重置")
        else:
            messages.append(f"累计计数首末差：{counters[-1]-counters[0]} 个通道失败（非无效帧数；首条计数不算本文件新增）")
    unavailable = sum(str(row.get("acquisition_diag_available", "")).strip() == "0"
                      for row in rows if row.get("record_type", "").strip().upper() == "PAIRED")
    if unavailable:
        messages.append(f"{unavailable} 条配对记录无固件诊断，无法据此确认是否发生采集错误")
    return {"invalid_cycles": invalid_cycles, "event_count": sum(events.values()),
            "invalid_paired": invalid_paired, "events": dict(events),
            "has_warning": bool(events or unavailable or (counters and max(counters) > 0)),
            "text": "；".join(messages)}


def csv_paths(source):
    root = Path(str(source).strip().strip('"'))
    if root.is_file():
        if root.suffix.lower() != ".csv":
            raise ValueError("请选择 CSV 文件或包含 CSV 的目录。")
        return [root]
    if not root.is_dir():
        raise ValueError(f"数据路径不存在：{root}")
    paths = sorted(p for p in root.rglob("*") if p.is_file() and p.suffix.lower() == ".csv")
    if not paths:
        raise ValueError(f"所选目录及子目录没有 CSV：{root}")
    return paths


def read_csv(path):
    data = path.read_bytes()
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("gb18030")
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return [], [], data
    header = [value.strip() for value in rows[0]]
    if len(set(header)) != len(header):
        raise ImportProblem("format", "CSV 存在重复列名")
    records = []
    for line, values in enumerate(rows[1:], 2):
        if not any(value.strip() for value in values):
            continue
        if len(values) != len(header):
            raise ImportProblem("format", f"第 {line} 行字段数 {len(values)} 与表头 {len(header)} 不符")
        record = dict(zip(header, values))
        record["_source_line"] = line
        records.append(record)
    return header, records, data


def _mark(text):
    found = []
    for match in MARK_PATTERN.finditer(unicodedata.normalize("NFKC", text)):
        position, force, trial = match.groups()
        position = "前" if position.lower() in {"前", "front"} else "后"
        found.append((position, float(force), int(trial) if trial else None))
    return found


def _condition(path, rows):
    # Recorded marks take precedence; never silently resolve conflicting labels.
    marks = {str(row.get("position_mark", "")).strip() for row in rows} - {""}
    parsed = set()
    for mark in marks:
        matches = _mark(mark)
        if len(matches) != 1:
            raise ImportProblem("mark", f"位置标记无法识别：{mark}；请使用 前15N_01 这样的标记")
        parsed.add(matches[0])
    filename = _mark(path.stem)
    if not parsed:
        parsed = set(filename)
    if len(parsed) != 1:
        raise ImportProblem("mixed_condition" if parsed else "mark", "一个 CSV 必须只含一次采集；未找到唯一的 前/后+力值N+次数 标记")
    condition = next(iter(parsed))
    if filename and any(value[:2] != condition[:2] for value in filename):
        raise ImportProblem("conflict", "文件名与 CSV 内的位置/力值标记冲突")
    if condition[:2] not in CONDITIONS:
        raise ImportProblem("condition", "该文件不属于前/后 3、15、25 N 六个工况")
    warning = ""
    if filename and any(value[2] is not None and condition[2] is not None and value[2] != condition[2] for value in filename):
        warning = "文件名与记录中的次数编号不同，工况一致；按此文件统计一次，请核对编号"
    return (*condition, warning)


def load_recordings(source, progress=None, *, allow_empty=False):
    paths = csv_paths(source)
    trials, sources = [], []
    seen_hashes, seen_frames, seen_trials = {}, {}, {}
    for index, path in enumerate(paths, 1):
        if progress:
            progress(5 + int(index / len(paths) * 55), f"读取 {path.name}")
        source_row = {"path": str(path), "status": "未导入", "detail": "", "rows": 0,
                      "reason_code": "", "reason": "", "action": "", "condition": "", "warnings": "",
                      "invalid_cycles": 0, "acquisition_events": 0, "diagnostics": ""}
        sources.append(source_row)
        try:
            header, rows, data = read_csv(path)
            source_row["rows"] = len(rows)
            diagnostics = acquisition_diagnostics(rows)
            source_row.update(invalid_cycles=diagnostics["invalid_cycles"],
                              acquisition_events=diagnostics["event_count"], diagnostics=diagnostics["text"])
            if not rows:
                raise ImportProblem("empty", "只有表头或内容为空，没有采样记录")
            missing = REQUIRED_COLUMNS - set(header)
            if missing:
                raise ImportProblem("format", "不是最新版“记录与导出”CSV；缺少列：" + ", ".join(sorted(missing)))
            continuous_reason = continuous_protocol_reason(rows)
            if continuous_reason:
                raise ImportProblem("continuous_protocol", continuous_reason)
            digest = hashlib.sha256(data).hexdigest()
            if digest in seen_hashes:
                raise ImportProblem("duplicate", f"内容与 {seen_hashes[digest]} 完全重复，未重复计数")
            position, force, trial_number, mark_warning = _condition(path, rows)
            source_row["condition"] = f"{position}区 {force:g} N"
            frames = [row for row in rows if row["record_type"].strip().upper() == "PAIRED"]
            if not frames:
                raise ImportProblem("no_frames", "record_type 列没有 PAIRED 采样行")
            tare_epochs = {row.get("host_tare_epoch", "").strip() for row in frames} - {""}
            if len(tare_epochs) > 1:
                raise ImportProblem("mode", "同一文件的 PAIRED 采样跨多个上位机置零基线（host_tare_epoch）；请按基线分别导出，不能混合评价")
            try:
                working_zero_notice = audit_effective_baseline(frames)
            except ValueError as error:
                raise ImportProblem("mode", str(error)) from error
            modes = {row["selected_value_mode"].strip() for row in frames}
            if len(modes) != 1 or "" in modes:
                raise ImportProblem("mode", "同一个文件缺少力值模式，或切换了力值模式")
            model_states = {(row.get("scalar_model_status", "").strip(), "HIGH_FORCE_CORRECTION_OFF" in _flags(row)) for row in frames}
            if len(model_states) != 1:
                raise ImportProblem("mode", "同一个文件存在不同模型状态或高力修正设置")
            device_ids = {row.get("device_sn", "").strip() for row in rows} - {""}
            if len(device_ids) > 1:
                raise ImportProblem("device", "同一个文件包含多个设备编号")
            flagged = any(row.get("stable_3s_sample", "").strip() in {"1", "1.0", "true", "True"} for row in frames)
            included, excluded, warnings = [], Counter(), Counter()
            for frame in frames:
                if invalid_paired_capture(frame):
                    excluded["采集诊断无效"] += 1
                elif flagged and frame.get("stable_3s_sample", "").strip() not in {"1", "1.0", "true", "True"}:
                    excluded["非稳定采样标记"] += 1
                elif frame.get("frame_state", "PAIRED").strip().upper() != "PAIRED":
                    excluded["失效帧"] += 1
                elif _flags(frame) - INFORMATION_FLAGS - MEASUREMENT_WARNINGS:
                    excluded["采样无效（" + "、".join(sorted(_flags(frame) - INFORMATION_FLAGS - MEASUREMENT_WARNINGS)) + "）"] += 1
                elif _finite(frame.get(VALUE_COLUMN)) is None:
                    excluded["力值无效"] += 1
                else:
                    included.append(frame)
                    warnings.update(_flags(frame) & MEASUREMENT_WARNINGS)
            if len(included) < 2:
                reasons = "；".join(f"{key} {count} 点" for key, count in excluded.items())
                raise ImportProblem("too_few", f"共 {len(frames)} 个采样点，可用 {len(included)} 个，不足 2 个。{reasons}")
            identities = {(row.get("local_time_iso", ""), row.get("device_ms", "")) for row in frames
                          if row.get("local_time_iso") and row.get("device_ms")}
            overlap = identities & seen_frames.keys()
            if overlap:
                previous = seen_frames[next(iter(overlap))]
                raise ImportProblem("overlap", f"与 {previous} 有 {len(overlap)} 个重复采样帧，疑似未清空本轮后再次导出；未重复计数")
            trial_key = (position, force, trial_number)
            if trial_number is not None and trial_key in seen_trials:
                mark_warning += ("；" if mark_warning else "") + f"次数编号与 {seen_trials[trial_key]} 相同；未发现相同采样，按文件分别统计，请核对编号"
            values = [float(row[VALUE_COLUMN]) for row in included]
            references = [_finite(row.get("measured_reference_resultant_n")) for row in included]
            reference = median(references) if all(value is not None and value > 0 for value in references) else None
            value = median(values)
            note = "；".join(f"{key} {count} 帧" for key, count in excluded.items())
            notices = set().union(*(_flags(frame) & (INFORMATION_FLAGS - {"OK"}) for frame in included))
            if notices:
                note += ("；" if note else "") + "；".join(FLAG_LABELS[flag] for flag in sorted(notices))
            warning_text = "；".join(f"{FLAG_LABELS[flag]}（{count} 点）" for flag, count in sorted(warnings.items()))
            if mark_warning:
                warning_text += ("；" if warning_text else "") + mark_warning
            if working_zero_notice:
                warning_text += ("；" if warning_text else "") + working_zero_notice
            if diagnostics["has_warning"]:
                warning_text += ("；" if warning_text else "") + diagnostics["text"]
            if warning_text:
                note += ("；" if note else "") + warning_text
            if diagnostics["text"] and not diagnostics["has_warning"]:
                note += ("；" if note else "") + diagnostics["text"]
            if not flagged:
                note = "无稳定区间标记，使用整段有效记录；请确认录制时已稳态" + ("；" + note if note else "")
            trials.append({
                "path": str(path), "position": position, "force": force, "trial": trial_number,
                "frames": len(frames), "used_frames": len(values), "median": value,
                "mean": mean(values), "std": stdev(values), "min": min(values), "max": max(values),
                "reference": reference, "measured_error": value - reference if reference is not None else None,
                "nominal_deviation": value - force, "mode": next(iter(modes)),
                "model_state": next(iter(model_states)),
                "device_sn": next(iter(device_ids), ""),
                "selection": "稳定采样标记帧" if flagged else "整段有效记录（待确认稳态）", "note": note,
                "warning_frames": sum(bool(_flags(frame) & MEASUREMENT_WARNINGS) for frame in included),
                "warnings": warning_text,
                "invalid_cycles": diagnostics["invalid_cycles"], "acquisition_events": diagnostics["event_count"],
                "diagnostics": diagnostics["text"],
            })
            source_row.update(status="已导入", detail=note, warnings=warning_text,
                              reason="已统计，含提示" if warning_text else "已统计",
                              action="请核对采集异常；结果仅覆盖有效记录，不能据此确认完整复测。" if diagnostics["has_warning"] else
                                     ("请核对测量提示，读数已保留。" if warning_text else "无需处理"))
            seen_hashes[digest] = path.name
            seen_frames.update({identity: path.name for identity in identities})
            if trial_number is not None:
                seen_trials[trial_key] = path.name
        except (ValueError, UnicodeError, OSError) as exc:
            source_row["detail"] = str(exc) + ("；" + source_row["diagnostics"] if source_row["diagnostics"] else "")
            code = exc.code if isinstance(exc, ImportProblem) else "read"
            source_row.update(reason_code=code, reason=PROBLEMS[code][0], action=PROBLEMS[code][1])
    if not trials and not allow_empty:
        details = "\n".join(f"{Path(row['path']).name}：{row['detail']}" for row in sources[:4])
        raise ValueError("没有可用于六工况评价的记录 CSV。无需 summary；请选择“记录与导出”生成的文件。\n" + details)
    if len({(trial["mode"], trial["model_state"]) for trial in trials}) > 1:
        raise ValueError("所选记录使用了不同力值模式、模型状态或高力修正设置，不能合并比较；请按相同设置分目录生成。")
    if len({trial["device_sn"] for trial in trials} - {""}) > 1:
        raise ValueError("六工况重复采集报告每次处理一台设备；请按设备分目录生成。")
    return trials, sources


def recording_groups(trials):
    groups = []
    for position, force in CONDITIONS:
        values = [trial["median"] for trial in trials if (trial["position"], trial["force"]) == (position, force)]
        count = len(values)
        average = mean(values) if values else None
        deviation = stdev(values) if count >= 2 else None
        groups.append({
            "position": position, "force": force, "count": count,
            "status": "数量齐全" if count == EXPECTED_TRIALS else (f"缺 {EXPECTED_TRIALS - count} 次" if count < EXPECTED_TRIALS else f"多 {count - EXPECTED_TRIALS} 次，请核对"),
            "mean": average, "std": deviation,
            "cv": deviation / abs(average) if deviation is not None and average != 0 else None,
            "range": max(values) - min(values) if count >= 2 else None,
            "min": min(values) if values else None, "max": max(values) if values else None,
            "warning_trials": sum(bool(t.get("warnings")) for t in trials if (t["position"], t["force"]) == (position, force)),
            "nominal_deviation": average - force if average is not None else None,
        })
    return groups


def generate_recording_report(source, output_path, progress=None):
    # Reuse the monitor's existing Python workbook writer so the EXE stays standalone.
    from consistency_report import _sheet_title, _write_table, _fit_columns

    trials, sources = load_recordings(source, progress)
    groups = recording_groups(trials)
    book = Workbook()
    summary = book.active
    summary.title = "重复采集评价"
    _sheet_title(summary, "前后六工况重复采集评价", "每个 CSV 算一次采集；先取文件内有效力值中位数，再计算组内重复性。每个工况预期 10 次。", 10)
    summary["A4"] = "当前力值模式"
    summary.merge_cells("B4:J4")
    summary["B4"] = trials[0]["mode"]
    summary.merge_cells("A5:J5")
    summary["A5"] = "本报告按同一台设备统计。CSV 未记录设备编号时，请勿混入其他设备的数据。"
    summary.merge_cells("A6:J6")
    summary["A6"] = "标称值来自位置标记或文件名；相对标称值偏差不等同于测力计实测误差。"
    summary.merge_cells("A7:J7")
    summary["A7"] = "本方案评价六个工况的重复性；不判定低力、迟滞、卸载残余或设备间一致性。"
    headers = ["位置", "标称力(N)", "有效文件数", "采集数量", "文件中位数的均值(N)", "文件间标准差(N)", "CV", "文件间极差(N)", "相对标称值偏差(N)", "预期次数"]
    keys = ("position", "force", "count", "status", "mean", "std", "cv", "range", "nominal_deviation")
    _write_table(summary, 9, headers, [[*[g[key] for key in keys], EXPECTED_TRIALS] for g in groups], "RecordingGroups")
    for row in range(10, 16):
        for col in (5, 6, 8, 9):
            summary.cell(row, col).number_format = "0.000"
        summary.cell(row, 7).number_format = "0.0%"
    summary_notes = [
        "未设置合格阈值，数量齐全仅表示每组恰好 10 个有效文件。",
        "无稳定采样标记的手动记录使用整段有效帧，不自动截取看起来最好的片段。",
        "采集请先施力并保持稳定，再开始记录；停止记录后卸载，每次导出后清空本轮。",
        f"已导入 {len(trials)} 个文件；未导入 {len(sources) - len(trials)} 个，原因见“导入说明”。",
    ]
    capture_events = sum(source["acquisition_events"] for source in sources)
    if capture_events:
        summary_notes.append(f"记录含 {capture_events} 条采集异常事件，其中 {sum(s['invalid_cycles'] for s in sources)} 个明确无效采集周期；事件不作 0 或旧值计入，结果仅覆盖有效记录。详见导入说明，不能按文件数齐全认定完整复测。")
    for row, note in enumerate(summary_notes, 18):
        summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        summary.cell(row, 1, note).alignment = Alignment(wrap_text=True, vertical="center")
        summary.row_dimensions[row].height = 30

    detail = book.create_sheet("每次采集")
    detail_headers = ["位置", "标称力(N)", "次数标记", "有效帧数", "配对帧数", "中位数(N)", "均值(N)", "帧内标准差(N)", "最小值(N)", "最大值(N)", "测力计参考中位数(N)", "相对实测值偏差(N)", "相对标称值偏差(N)", "采样范围", "说明", "来源CSV"]
    _sheet_title(detail, "每次采集明细", "一行对应一个 CSV。帧内标准差描述单次波动，文件间标准差描述重复采集差异。空白参考值表示没有有效实测参考。", len(detail_headers))
    keys = ("position", "force", "trial", "used_frames", "frames", "median", "mean", "std", "min", "max", "reference", "measured_error", "nominal_deviation", "selection", "note", "path")
    ordered = sorted(trials, key=lambda t: (CONDITIONS.index((t["position"], t["force"])), t["trial"] or 0, t["path"]))
    _write_table(detail, 4, detail_headers, [[trial[key] for key in keys] for trial in ordered], "RecordingTrials")
    for row in detail.iter_rows(min_row=5, min_col=6, max_col=13):
        for cell in row:
            cell.number_format = "0.000"
    imported = book.create_sheet("导入说明")
    _sheet_title(imported, "文件导入结果", "原始文件只读。空文件、不支持的格式、重复文件和不明确的工况均列出原因。", 4)
    _write_table(imported, 4, ["CSV路径", "数据行数", "状态", "说明"], [[row[key] for key in ("path", "rows", "status", "detail")] for row in sources], "RecordingSources")
    for sheet in book:
        sheet.sheet_view.showGridLines = False
        _fit_columns(sheet, 28)
        sheet.row_dimensions[9 if sheet is summary else 4].height = 42
    for column in ("A", "B", "C", "D", "J"):
        summary.column_dimensions[column].width = 16 if column != "D" else 26
    for column in ("E", "F", "H", "I"):
        summary.column_dimensions[column].width = 23
    detail.column_dimensions["N"].width = 32
    detail.column_dimensions["O"].width = 70
    detail.column_dimensions["P"].width = 80
    imported.column_dimensions["A"].width = 85
    imported.column_dimensions["D"].width = 85
    for sheet in (detail, imported):
        for row in range(5, sheet.max_row + 1):
            sheet.row_dimensions[row].height = 44
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    if progress:
        progress(100, "六工况重复采集报告已生成")
    return {"path": str(path), "report_mode": "recordings", "file_count": len(trials),
            "skipped_files": len(sources) - len(trials),
            "complete_groups": sum(g["count"] == EXPECTED_TRIALS for g in groups)}
